"""
Independent verification of PCL_05: reads the VALUES the recalculated workbook displays
(openpyxl data_only) and compares them with the SQL export CSVs - no workbook formula is reused.
Usage (from the repository root): python3 qa/verify_against_sql.py data/excel_import excel/PCL_05_Porsche_Margin_Model_FINAL.xlsx
"""
import csv, json, os, sys
import openpyxl

X, WB = sys.argv[1], sys.argv[2]
M = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "PCL_05_cell_map.json")))   # cell map written by tools/build_workbook.py
wb = openpyxl.load_workbook(WB, data_only=True)
rd = lambda f: list(csv.reader(open(os.path.join(X, f), encoding="utf-8")))
P, PC = M["P"], M["PC"]
res = []   # (area, item, n_compared, n_fail, max_abs_diff)

def num(v):
    return None if v in (None, "") else float(v)

def cmp(area, item, pairs, tol):
    fails, mx = 0, 0.0
    for a, b in pairs:
        if a is None and b is None: continue
        if a is None or b is None: fails += 1; continue
        d = abs(a - b); mx = max(mx, d); fails += d > tol
    res.append((area, item, len(pairs), fails, mx, tol))

# 1. Data_Panel grid, cell by cell, against core.fact_financial export (value AND blank-ness)
fl = {(r[0], r[1]): float(r[3]) for r in rd("fact_long.csv")[1:]}
dp = wb["Data_Panel"]
pairs = []
for v, rr in M["PROW"].items():
    for p, c in zip(P, PC):
        pairs.append((num(dp[f"{c}{rr}"].value), fl.get((p, v))))
cmp("1 Data_Panel", f"{len(M['PROW'])} variables x 7 periods vs core.fact_financial", pairs, 1e-6)

# 2. Analysis-sheet rows vs SQL mart views (sqlout.csv), displayed values
so = rd("sqlout.csv"); h = so[0]; S = {r[0]: r for r in so[1:]}
col_of = {}  # SQL_Outputs row -> sql column name
sq = wb["SQL_Outputs"]
for r in range(1, 40):
    if sq.cell(r, 1).value in h: col_of[r] = sq.cell(r, 1).value
for cid, xl, sqlr, mult in M["recon"]:
    sh, rng = xl.split("!"); c1, c2 = rng.split(":")
    row = int("".join(ch for ch in c1 if ch.isdigit())); first = c1.rstrip("0123456789")
    srow = int(sqlr.split("!")[1].split(":")[0][1:])
    name = col_of.get(srow) or sq.cell(srow, 2).value
    dec = max((len(S[p][h.index(name)].split(".")[1]) if "." in S[p][h.index(name)] else 0) for p in P if S[p][h.index(name)] != "")
    pairs = []
    for p, c in zip(P, PC):
        if c < first: continue
        xv = num(wb[sh][f"{c}{row}"].value); sv = num(S[p][h.index(name)])
        pairs.append((None if xv is None else round(xv * mult, dec), sv))
    cmp("2 Analysis vs SQL", f"{sh} row {row} = {name}", pairs, 0.5 * 10 ** -dec + 1e-9)

# 3. KPI_Overview displayed rows vs SQL (headline figures)
k = wb["KPI_Overview"]
lab = {k.cell(r, 2).value: r for r in range(1, 40) if k.cell(r, 2).value}
KMAP = [("L0 Reported EBIT", "ebit_l0_reported", 1), ("L0 Reported margin", "margin_l0_reported_pct", 100),
        ("L1 EBIT excl. realignment & battery", "ebit_l1_ex_realignment", 1), ("L1 margin", "margin_l1_pct", 100),
        ("L2 EBIT excl. realignment, battery & tariffs", "ebit_l2_ex_realignment_tariffs", 1), ("L2 margin", "margin_l2_pct", 100),
        ("ASP (automotive revenue per vehicle sold)", "asp_calc_eur_k", 1), ("Net cash flow margin", "ncf_margin_pct", 100)]
for label, name, mult in KMAP:
    rr = lab[label]; dec = 6 if mult == 1 and "ebit" in name else 2
    pairs = [(round(num(k[f"{c}{rr}"].value) * mult, dec), num(S[p][h.index(name)])) for p, c in zip(P, PC)]
    cmp("3 KPI_Overview vs SQL", f"{label}", pairs, 0.5 * 10 ** -dec + 1e-9)
rr = lab["Group sales revenue"]
cmp("3 KPI_Overview vs SQL", "Group sales revenue (G01)", [(num(k[f"{c}{rr}"].value), fl[(p, "G01")]) for p, c in zip(P, PC)], 1e-6)

# 4. Guidance scenarios vs v_guidance_2026 / _summary / _auto
g = wb["Guidance_Scenarios"]
grid = {g.cell(r, 1).value: r for r in range(49, 76)}
gg = rd("ggrid.csv"); gh = gg[0]
pairs = []
for r in gg[1:]:
    rr = grid[r[0]]
    for col, name in (("L", "h2_margin_l0_pct"), ("O", "h2_margin_l1_pct"), ("P", "h2_margin_l2_pct")):
        pairs.append((round(g[f"{col}{rr}"].value * 100, 2), float(r[gh.index(name)])))
cmp("4 Guidance vs SQL", "27 combinations x L0/L1/L2 H2 margin", pairs, 0.005 + 1e-9)
gs = rd("gsum.csv"); sh_ = gs[0]; s1 = gs[1]
pairs = []
for rr, lvl in ((79, "l0"), (80, "l1"), (81, "l2")):
    for col, suf in (("C", "min"), ("D", "mid"), ("E", "max")):
        pairs.append((round(g[f"{col}{rr}"].value * 100, 2), float(s1[sh_.index(f"h2_margin_{lvl}_{suf}")])))
cmp("4 Guidance vs SQL", "Summary min / all-midpoints / max (L0, L1, L2)", pairs, 0.005 + 1e-9)
ga = rd("gauto.csv"); ah = ga[0]; auto = {g.cell(r, 1).value: r for r in range(88, 97)}
pairs = []
for r in ga[1:]:
    rr = auto[r[0]]
    pairs += [(round(g[f"K{rr}"].value * 100, 2), float(r[1])), (round(g[f"M{rr}"].value * 100, 2), float(r[2]))]
cmp("4 Guidance vs SQL", "9 automotive combinations (H2 EBITDA and NCF margin)", pairs, 0.005 + 1e-9)

# 5. Formula errors anywhere
errs = [(ws.title, c.coordinate, c.value) for ws in wb for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("#") and c.value.rstrip("!?/0").upper() in ("#REF", "#NAME", "#VALUE", "#DIV", "#N/A", "#NUM", "#NULL")]
res.append(("5 Error scan", "cells showing an Excel error value", sum(1 for _ in wb.worksheets), len(errs), 0.0, 0))

print(f"{'area':22} {'item':70} {'n':>4} {'fail':>4} {'max|diff|':>11}")
for a, i, n, f, mx, tol in res:
    print(f"{a:22} {i[:70]:70} {n:>4} {f:>4} {mx:>11.2e}")
tf = sum(r[3] for r in res)
print(f"\nTOTAL comparisons: {sum(r[2] for r in res[:-1])}   failures: {tf}   error cells: {len(errs)}")
print("RESULT:", "PASS" if tf == 0 else "FAIL")
