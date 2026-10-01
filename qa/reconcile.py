"""
PCL Phase 6 - reconciliation of the Power BI layer to the approved SQL and Excel model.

Power BI Desktop cannot be driven from a script, so every DAX measure used on the dashboard or the QA page
is replicated here line for line (same formula, same filter logic) over the SAME CSV tables Power BI loads.
The replica is then reconciled to:
  R1  RefSQL.csv        - what the in-Power BI QA page will compare against (expected: all PASS)
  R2  live SQL views    - queried directly from PostgreSQL (proves the CSV export is faithful)
  R3  PCL_05 workbook   - the approved Excel model (cached values after full recalculation)
  R4  identities        - bridge closure, pp closure, volume+ASP, mix totals, NCF identity, L0->L1->L2 waterfall
  R5  export integrity  - every base value in FactHalfYear equals mart.v_panel_hy
Usage (from the repository root): python3 qa/reconcile.py excel/PCL_05_Porsche_Margin_Model_FINAL.xlsx
Needs the PostgreSQL database built by sql/00_run_all.sql (for R2).
"""
import csv, json, os, subprocess, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
D = os.path.join(ROOT, "data", "powerbi")          # the CSV tables the .pbix loads
def load(t):
    return list(csv.DictReader(open(os.path.join(D, t + ".csv"), encoding="utf-8")))
def f(x):
    return None if x in ("", None) else float(x)

PER = load("DimPeriod"); FH = {r["period_code"]: r for r in load("FactHalfYear")}
BR = load("FactBridgeYoY"); SEN = load("FactSensitivity"); MIX = load("FactDeliveryMix")
GD = {r["scenario_id"]: r for r in load("FactGuidance")}; GA = {r["scenario_id"]: r for r in load("FactGuidanceAuto")}
REF = load("RefSQL")
P = [r["period_code"] for r in sorted(PER, key=lambda r: int(r["sort_key"]))]
PRIOR = {r["period_code"]: (r["prior_same_half_code"] or None) for r in PER}
LATEST = [r["period_code"] for r in PER if r["is_latest"] == "true"][0]
FY = {r["period_code"]: int(r["fiscal_year"]) for r in PER}; HALF = {r["period_code"]: int(r["half_no"]) for r in PER}

# ------------------------------------------------------------------ DAX replica (context = one period p)
def S(col, p): return f(FH[p][col]) or 0.0 if FH[p][col] != "" else None   # SUM over one row; blank -> None
def s0(col, p):
    v = f(FH[p][col]); return 0.0 if v is None else v
def div(a, b): return None if a is None or b in (None, 0) else a / b
def nb(*xs): return all(x is not None for x in xs)

class R:
    # 01 levels
    Revenue = staticmethod(lambda p: s0("grp_revenue", p))
    L0 = staticmethod(lambda p: s0("grp_ebit", p))
    AdjR = staticmethod(lambda p: s0("adj_realignment_net", p))
    AdjT = staticmethod(lambda p: s0("adj_us_tariffs", p))
    L1 = staticmethod(lambda p: R.L0(p) + R.AdjR(p))
    L2 = staticmethod(lambda p: R.L1(p) + R.AdjT(p))
    M0 = staticmethod(lambda p: div(R.L0(p), R.Revenue(p)))
    M1 = staticmethod(lambda p: div(R.L1(p), R.Revenue(p)))
    M2 = staticmethod(lambda p: div(R.L2(p), R.Revenue(p)))
    Band = staticmethod(lambda p: s0("rounding_band_eur_m", p))
    # 02 L3
    CleanDA = staticmethod(lambda p: s0("da_total", p) - s0("impairment_used", p))
    CapDev = staticmethod(lambda p: s0("rd_capitalised", p))
    L3 = staticmethod(lambda p: R.L1(p) + R.CleanDA(p) - R.CapDev(p))
    M3 = staticmethod(lambda p: div(R.L3(p), R.Revenue(p)))
    # 05 R&D
    RDT = staticmethod(lambda p: s0("rd_total", p))
    RDE = staticmethod(lambda p: s0("rd_expensed", p))
    AMO = staticmethod(lambda p: s0("amort_capitalised_rd", p))
    RDPL = staticmethod(lambda p: R.RDE(p) + R.AMO(p))
    CapRate = staticmethod(lambda p: div(R.CapDev(p), R.RDT(p)))
    NetCap = staticmethod(lambda p: R.CapDev(p) - R.AMO(p))
    # 06 commercial
    AutoRev = staticmethod(lambda p: s0("auto_revenue", p))
    VS = staticmethod(lambda p: s0("vehicle_sales", p))
    Deliv = staticmethod(lambda p: s0("deliveries", p))
    ASP = staticmethod(lambda p: div(R.AutoRev(p) * 1000, R.VS(p)))
    # 07 cash
    NCF = staticmethod(lambda p: s0("auto_net_cash_flow", p))
    CFO = staticmethod(lambda p: s0("auto_cfo", p))
    CFI = staticmethod(lambda p: s0("auto_cfi_operating", p))
    EBITDA = staticmethod(lambda p: s0("auto_ebitda", p))
    AEBIT = staticmethod(lambda p: s0("auto_ebit", p))
    Capex = staticmethod(lambda p: s0("capex", p))
    NCFM = staticmethod(lambda p: div(R.NCF(p), R.AutoRev(p)))
    EBM = staticmethod(lambda p: div(R.EBITDA(p), R.AutoRev(p)))
    CFOE = staticmethod(lambda p: div(R.CFO(p), R.EBITDA(p)))
    NCFE = staticmethod(lambda p: div(R.NCF(p), R.AEBIT(p)) if R.AEBIT(p) > 0 else None)
    Proxy = staticmethod(lambda p: R.EBITDA(p) - R.Capex(p) - R.CapDev(p))
    FS = staticmethod(lambda p: s0("fs_ebit", p))

def PY(fn, p):
    q = PRIOR[p]; return fn(q) if q else None
def YOY(fn, p, kind="diff"):
    py = PY(fn, p); c = fn(p)
    if py is None or c is None: return None
    return {"diff": c - py, "pp": (c - py) * 100, "pct": div(c, py) - 1 if div(c, py) is not None else None}[kind]
def capfx(p):   # Cap Rate Effect (DAX recompute)
    cpy = PY(R.CapRate, p)
    return None if cpy is None else R.RDPL(p) - (R.RDT(p) * (1 - cpy) + R.AMO(p))
def volfx(p):
    vpy, apy = PY(R.VS, p), PY(R.ASP, p)
    return None if vpy is None else (R.VS(p) - vpy) * apy / 1000
def aspfx(p):
    apy = PY(R.ASP, p)
    return None if apy is None else (R.ASP(p) - apy) * R.VS(p) / 1000
def rec_eff(p):
    if not PRIOR[p]: return None
    q = PRIOR[p]
    e = {"exc": -(R.AdjR(p) - R.AdjR(q)), "tar": -(R.AdjT(p) - R.AdjT(q)), "cap": R.CapDev(p) - R.CapDev(q),
         "da": -(R.CleanDA(p) - R.CleanDA(q)), "fs": R.FS(p) - R.FS(q)}
    e["res"] = YOY(R.L0, p) - sum(e.values()); return e
def bridge(p, step, col="eur_m"):
    v = [f(r[col]) for r in BR if r["period_code"] == p and r["step_id"] == step]
    return v[0] if v else None
def bridge_total(p, col="eur_m"):
    v = [f(r[col]) for r in BR if r["period_code"] == p and r["step_id"] != "S0_START" and r[col] != ""]
    return sum(v) if v else None
def sens(p, sid):
    v = [f(r["value_eur_m"]) for r in SEN if r["period_code"] == p and r["sensitivity_id"] == sid]; return v[0] if v else None
# guidance
def g_sel(sid, col): return f(GD[sid][col]) / 100
def g_min(col): return min(f(r[col]) for r in GD.values()) / 100
def g_max(col): return max(f(r[col]) for r in GD.values()) / 100
def g_mid(col): return f(GD["mid/mid/mid"][col]) / 100
def g_recalc(sid, lvl):
    r = f(GD[sid]["fy_rev"]); ros = f(GD[sid]["fy_ros"]); exc = f(GD[sid]["fy_exc"]); L = LATEST
    num = r * ros / 100 - R.L0(L)
    if lvl >= 1: num += exc - R.AdjR(L)
    if lvl >= 2: num += R.AdjT(L)
    return num / (r - R.Revenue(L))
def ga_recalc(sid, which):
    L = LATEST; share = R.AutoRev(L) / R.Revenue(L); fya = f(GA[sid]["fy_grp_rev"]) * share
    col, h1 = ("fy_ebitda_margin", R.EBITDA(L)) if which == "ebitda" else ("fy_ncf_margin", R.NCF(L))
    return (f(GA[sid][col]) / 100 * fya - h1) / (fya - R.AutoRev(L))
def mix_share(p, dim, mid):
    rows = [r for r in MIX if r["period_code"] == p and r["dimension"] == dim]
    tot = sum(f(r["deliveries"]) for r in rows); v = [f(r["deliveries"]) for r in rows if r["member_id"] == mid]
    return v[0] / tot if v else None

QA = {  # metric_key -> replica of the QA DAX Value expression (per period)
 "ebit_l0_reported": R.L0, "ebit_l1_ex_realignment": R.L1, "ebit_l2_ex_realignment_tariffs": R.L2, "ebit_l3_hybrid_sensitivity": R.L3,
 "margin_l0_reported_pct": lambda p: R.M0(p) * 100, "margin_l1_pct": lambda p: R.M1(p) * 100, "margin_l2_pct": lambda p: R.M2(p) * 100,
 "margin_l3_hybrid_pct": lambda p: R.M3(p) * 100, "yoy_ebit_l1": lambda p: YOY(R.L1, p), "yoy_margin_l1_pp": lambda p: YOY(R.M1, p, "pp"),
 "yoy_margin_l0_pp": lambda p: YOY(R.M0, p, "pp"), "da_clean": R.CleanDA, "yoy_da_clean": lambda p: YOY(R.CleanDA, p),
 "rd_total_costs": R.RDT, "rd_pl_charge": R.RDPL, "cap_rate_calc_pct": lambda p: R.CapRate(p) * 100, "net_capitalisation": R.NetCap,
 "ebit_effect_of_cap_rate_change": capfx, "yoy_rd_pl_charge": lambda p: YOY(R.RDPL, p), "asp_calc_eur_k": R.ASP,
 "yoy_vehicle_sales_pct": lambda p: None if YOY(R.VS, p, "pct") is None else YOY(R.VS, p, "pct") * 100,
 "yoy_asp_pct": lambda p: None if YOY(R.ASP, p, "pct") is None else YOY(R.ASP, p, "pct") * 100,
 "yoy_auto_revenue": lambda p: YOY(R.AutoRev, p), "volume_effect": volfx, "asp_effect": aspfx, "auto_net_cash_flow": R.NCF,
 "ncf_margin_pct": lambda p: R.NCFM(p) * 100, "ebitda_margin_calc_pct": lambda p: R.EBM(p) * 100, "cfo_to_ebitda": R.CFOE,
 "simple_cash_proxy": R.Proxy, "d_ebit": lambda p: YOY(R.L0, p), "eff_residual_other_drivers": lambda p: rec_eff(p)["res"] if rec_eff(p) else None,
 "d_margin_pp": lambda p: YOY(R.M0, p, "pp"),
}
def qa_value(row):
    k, p, s = row["metric_key"], row["period_code"], row["scenario_id"]
    if k == "mix_share_pct": return mix_share(p, row["dimension"], row["member_id"]) * 100
    if k.startswith("gs_"):
        _, lv, st = k.split("_"); col = f"h2_margin_{lv}_pct"
        return {"min": g_min, "mid": g_mid, "max": g_max}[st](col) * 100
    if k.startswith("g_"): return g_recalc(s, int(k[len("g_h2_margin_l")])) * 100
    if k.startswith("ga_"): return ga_recalc(s, "ebitda" if "ebitda" in k else "ncf") * 100
    return QA[k](p)

RES = []
def rec(group, item, a, b, tol):
    ok = (a is None and b is None) or (a is not None and b is not None and abs(a - b) <= tol)
    RES.append((group, item, a, b, None if (a is None or b is None) else abs(a - b), tol, ok))

# ---- R1: replica of the in-Power BI QA page vs RefSQL
QROWS = []
for r in REF:
    _dv = qa_value(r); _sv = f(r["sql_value"])
    QROWS.append([r["metric_key"], r["scenario_id"] or r["period_code"], _dv, _sv, None if _dv is None else abs(_dv - _sv),
                  "PASS" if _dv is not None and abs(_dv - _sv) <= f(r["tolerance"]) else "FAIL", r["sql_view"]])
    rec("R1 QA page vs RefSQL", f"{r['metric_key']} {r['period_code']} {r['scenario_id']} {r.get('dimension','')} {r.get('member_id','')}", qa_value(r), f(r["sql_value"]), f(r["tolerance"]))

# ---- R2: live SQL views (psql) vs replica  (independent of the CSV export)
def psql(sql):
    out = subprocess.run(["psql", "-h", os.environ.get("PGHOST", "/tmp"), "-p", os.environ.get("PGPORT", "5433"), "-U", os.environ.get("PGUSER", "postgres"), "-d", os.environ.get("PGDATABASE", "porsche_lens"), "-At", "-F", "\t", "-c", sql],
                         capture_output=True, text=True, check=True).stdout
    return [l.split("\t") for l in out.strip().split("\n") if l]
LIVE = [("mart.v_ebit_adjusted", "ebit_l0_reported", R.L0, 1e-6), ("mart.v_ebit_adjusted", "ebit_l1_ex_realignment", R.L1, 1e-6),
        ("mart.v_ebit_adjusted", "ebit_l2_ex_realignment_tariffs", R.L2, 1e-6), ("mart.v_ebit_adjusted", "ebit_l3_hybrid_sensitivity", R.L3, 1e-6),
        ("mart.v_ebit_adjusted", "margin_l0_reported_pct", lambda p: R.M0(p) * 100, 0.005), ("mart.v_ebit_adjusted", "margin_l1_pct", lambda p: R.M1(p) * 100, 0.005),
        ("mart.v_ebit_adjusted", "margin_l2_pct", lambda p: R.M2(p) * 100, 0.005), ("mart.v_ebit_adjusted", "rounding_band_eur_m", R.Band, 1e-6),
        ("mart.v_da_clean", "da_clean", R.CleanDA, 1e-6), ("mart.v_rd_capitalisation", "rd_pl_charge", R.RDPL, 1e-6),
        ("mart.v_rd_capitalisation", "net_capitalisation", R.NetCap, 1e-6), ("mart.v_rd_capitalisation", "ebit_effect_of_cap_rate_change", capfx, 0.05),
        ("mart.v_asp_volume", "asp_calc_eur_k", R.ASP, 0.005), ("mart.v_asp_volume", "volume_effect", volfx, 0.05), ("mart.v_asp_volume", "asp_effect", aspfx, 0.05),
        ("mart.v_cash_crosscheck", "ncf_margin_pct", lambda p: R.NCFM(p) * 100, 0.005), ("mart.v_cash_crosscheck", "simple_cash_proxy", R.Proxy, 1e-6),
        ("mart.v_cash_crosscheck", "cfo_to_ebitda", R.CFOE, 0.005), ("mart.v_cash_crosscheck", "ncf_to_ebit", R.NCFE, 0.005)]
for view, col, fn, tol in LIVE:
    for pc, v in psql(f"select period_code, {col} from {view} order by sort_key"):
        rec("R2 replica vs live SQL", f"{view}.{col} {pc}", fn(pc), f(v), tol)
for pc, *vals in psql("select period_code, eff_exceptional_items, eff_us_tariffs, eff_rd_capitalisation, eff_clean_da, eff_financial_services, eff_residual_other_drivers from mart.v_ebit_bridge_yoy order by sort_key"):
    e = rec_eff(pc)
    for key, v in zip(["exc", "tar", "cap", "da", "fs", "res"], vals):
        rec("R2 replica vs live SQL", f"v_ebit_bridge_yoy {key} {pc} (DAX recompute)", e[key], f(v), 1e-6)
    for sid, v in zip(["S1_EXC", "S2_TAR", "S3_CAPDEV", "S4_DA", "S5_FS", "S6_RES"], vals):
        rec("R2 replica vs live SQL", f"FactBridgeYoY {sid} {pc} (imported)", bridge(pc, sid), f(v), 1e-6)
for sid, *vals in psql("select scenario_id, h2_margin_l0_pct, h2_margin_l1_pct, h2_margin_l2_pct from mart.v_guidance_2026"):
    for lv, v in enumerate(vals):
        rec("R2 replica vs live SQL", f"v_guidance_2026 L{lv} {sid} (DAX recompute)", g_recalc(sid, lv) * 100, f(v), 0.005)
for pc, dim, mid, share in psql("select period_code, dimension, member_id, share_pct from mart.v_delivery_mix"):
    rec("R2 replica vs live SQL", f"v_delivery_mix share {pc} {dim} {mid}", mix_share(pc, dim, mid) * 100, f(share), 0.005)

# ---- R3: approved Excel workbook (PCL_05) vs replica (Excel is unrounded -> tight tolerances)
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
COLS = "DEFGHIJ"
XL_NOTE = []
XL = [("EBIT_Analysis", 7, R.Revenue, 1e-6), ("EBIT_Analysis", 8, R.L0, 1e-6), ("EBIT_Analysis", 9, R.M0, 1e-9), ("EBIT_Analysis", 13, R.AdjR, 1e-6),
      ("EBIT_Analysis", 15, R.AdjT, 1e-6), ("EBIT_Analysis", 19, R.L1, 1e-6), ("EBIT_Analysis", 20, R.M1, 1e-9), ("EBIT_Analysis", 21, R.L2, 1e-6),
      ("EBIT_Analysis", 22, R.M2, 1e-9), ("EBIT_Analysis", 23, R.Band, 1e-6), ("EBIT_Analysis", 25, lambda p: YOY(R.L0, p), 1e-6),
      ("EBIT_Analysis", 26, lambda p: YOY(R.M0, p, "pp"), 1e-9), ("EBIT_Analysis", 27, lambda p: YOY(R.L1, p), 1e-6),
      ("EBIT_Analysis", 28, lambda p: YOY(R.M1, p, "pp"), 1e-9), ("EBIT_Analysis", 29, lambda p: YOY(R.L2, p), 1e-6),
      ("EBIT_Analysis", 30, lambda p: YOY(R.M2, p, "pp"), 1e-9), ("EBIT_Analysis", 33, R.CleanDA, 1e-6), ("EBIT_Analysis", 34, R.CapDev, 1e-6),
      ("EBIT_Analysis", 35, R.L3, 1e-6), ("EBIT_Analysis", 36, R.M3, 1e-9), ("EBIT_Analysis", 37, lambda p: sens(p, "SENS_A04_L3"), 1e-6),
      ("EBIT_Analysis", 39, lambda p: bridge(p, "S0_START"), 1e-6), ("EBIT_Analysis", 41, lambda p: bridge(p, "S1_EXC"), 1e-6),
      ("EBIT_Analysis", 42, lambda p: bridge(p, "S2_TAR"), 1e-6), ("EBIT_Analysis", 43, lambda p: bridge(p, "S3_CAPDEV"), 1e-6),
      ("EBIT_Analysis", 44, lambda p: bridge(p, "S4_DA"), 1e-6), ("EBIT_Analysis", 45, lambda p: bridge(p, "S5_FS"), 1e-6),
      ("EBIT_Analysis", 46, lambda p: bridge(p, "S6_RES"), 1e-6), ("EBIT_Analysis", 47, lambda p: sens(p, "SENS_A04_RES"), 1e-6),
      ("EBIT_Analysis", 50, lambda p: bridge(p, "S7_DENOM", "pp"), 0.0006), ("EBIT_Analysis", 51, lambda p: bridge(p, "S1_EXC", "pp"), 0.0006),
      ("EBIT_Analysis", 54, lambda p: bridge(p, "S4_DA", "pp"), 0.0006), ("EBIT_Analysis", 56, lambda p: bridge(p, "S6_RES", "pp"), 0.0006),
      ("RD_DA", 9, R.RDT, 1e-6), ("RD_DA", 10, R.CapDev, 1e-6), ("RD_DA", 11, R.RDE, 1e-6), ("RD_DA", 12, R.AMO, 1e-6), ("RD_DA", 13, R.RDPL, 1e-6),
      ("RD_DA", 14, R.CapRate, 1e-9), ("RD_DA", 17, R.NetCap, 1e-6), ("RD_DA", 21, capfx, 1e-6), ("RD_DA", 22, lambda p: YOY(R.RDPL, p), 1e-6),
      ("RD_DA", 28, lambda p: s0("da_total", p), 1e-6), ("RD_DA", 32, lambda p: s0("impairment_used", p), 1e-6), ("RD_DA", 34, R.CleanDA, 1e-6),
      ("Commercial", 7, R.AutoRev, 1e-6), ("Commercial", 8, R.VS, 1e-6), ("Commercial", 9, R.Deliv, 1e-6), ("Commercial", 11, R.ASP, 1e-9),
      ("Commercial", 17, lambda p: YOY(R.VS, p, "pct"), 1e-9), ("Commercial", 18, lambda p: YOY(R.ASP, p, "pct"), 1e-9),
      ("Commercial", 19, lambda p: YOY(R.AutoRev, p), 1e-6), ("Commercial", 20, volfx, 1e-6), ("Commercial", 21, aspfx, 1e-6),
      ("Cash_Crosscheck", 8, R.AEBIT, 1e-6), ("Cash_Crosscheck", 9, R.EBITDA, 1e-6), ("Cash_Crosscheck", 10, R.CFO, 1e-6), ("Cash_Crosscheck", 11, R.CFI, 1e-6),
      ("Cash_Crosscheck", 12, R.NCF, 1e-6), ("Cash_Crosscheck", 16, R.NCFM, 1e-9), ("Cash_Crosscheck", 19, R.EBM, 1e-9), ("Cash_Crosscheck", 22, R.CFOE, 1e-9),
      ("Cash_Crosscheck", 25, R.Capex, 1e-6), ("Cash_Crosscheck", 29, R.Proxy, 1e-6)]
for sh, row, fn, tol in XL:
    ws = wb[sh]; label = ws.cell(row, 2).value
    for p, c in zip(P, COLS):
        xv = ws[f"{c}{row}"].value; xv = None if xv in ("", None) else float(xv); dv = fn(p)
        if dv is None and xv == 0: xv = None           # Excel writes 0 where DAX returns BLANK (no prior year)
        if dv is None and xv is not None and PRIOR[p] is None:
            XL_NOTE.append(f"{sh}!{c}{row} {label.strip()[:45]} {p}: Excel shows {xv:,.2f}; Power BI BLANK (no prior-year half in window)")
            continue
        rec("R3 replica vs Excel PCL_05", f"{sh}!{c}{row} {label.strip()[:40]} {p}", dv, xv, tol)
gs = wb["Guidance_Scenarios"]
for rr, lv in ((79, "l0"), (80, "l1"), (81, "l2")):
    for c, fn in (("C", g_min), ("D", g_mid), ("E", g_max)):
        rec("R3 replica vs Excel PCL_05", f"Guidance_Scenarios!{c}{rr} {lv}", fn(f"h2_margin_{lv}_pct"), gs[f"{c}{rr}"].value, 0.00005)
grid = {gs.cell(r, 1).value: r for r in range(49, 76)}
for sid, r in grid.items():
    for lv, c in enumerate("LOP"):
        rec("R3 replica vs Excel PCL_05", f"Guidance grid {sid} L{lv} (DAX recompute)", g_recalc(sid, lv), gs[f"{c}{r}"].value, 1e-9)
kpi = wb["KPI_Overview"]
for label, fn in (("Group sales revenue", R.Revenue), ("L0 Reported EBIT", R.L0), ("L1 margin", R.M1), ("ASP (automotive revenue per vehicle sold)", R.ASP)):
    rr = [r for r in range(1, 40) if kpi.cell(r, 2).value == label][0]
    for p, c in zip(P, COLS):
        rec("R3 replica vs Excel PCL_05", f"KPI_Overview {label[:25]} {p}", fn(p), kpi[f"{c}{rr}"].value, 1e-9 if "margin" in label or "ASP" in label else 1e-6)

# ---- R4: identities that the dashboard relies on
for p in P:
    if PRIOR[p]:
        rec("R4 identities", f"Bridge closure EUR {p}", bridge_total(p), YOY(R.L0, p), 1e-6)
        rec("R4 identities", f"Bridge closure pp {p}", bridge_total(p, "pp"), YOY(R.M0, p, "pp"), 0.005)
        vol = f(FH[p]["sql_volume_effect"]); asp = f(FH[p]["sql_asp_effect"])
        rec("R4 identities", f"Volume + ASP effect = revenue change {p}", vol + asp, YOY(R.AutoRev, p), 0.1)
        rec("R4 identities", f"Revenue waterfall total (PY + volume + ASP) = automotive revenue {p}", PY(R.AutoRev, p) + vol + asp, R.AutoRev(p), 0.1)
    rec("R4 identities", f"L0->L1->L2 waterfall total = L2 {p}", R.L0(p) + R.AdjR(p) + R.AdjT(p), R.L2(p), 1e-9)
    rec("R4 identities", f"Mix model total = deliveries {p}", sum(f(r["deliveries"]) for r in MIX if r["period_code"] == p and r["dimension"] == "model"), R.Deliv(p), 0)
    rec("R4 identities", f"Mix region total = deliveries {p}", sum(f(r["deliveries"]) for r in MIX if r["period_code"] == p and r["dimension"] == "region"), R.Deliv(p), 0)
    rec("R4 identities", f"NCF = CFO + CFI {p}", R.CFO(p) + R.CFI(p), R.NCF(p), 0.5)
rec("R4 identities", "Guidance grid has 27 combinations", len(GD), 27, 0)
rec("R4 identities", "Automotive grid has 9 combinations", len(GA), 9, 0)
rec("R4 identities", "Every guidance row labelled not-a-forecast", sum("not a forecast" in r["scenario_type"] for r in list(GD.values()) + list(GA.values())), 36, 0)

# ---- R5: export integrity (FactHalfYear base columns = mart.v_panel_hy, cell by cell)
base = ["grp_revenue", "grp_ebit", "auto_revenue", "auto_ebit", "auto_ebitda", "rd_total", "rd_capitalised", "rd_expensed", "amort_capitalised_rd",
        "capex", "dep_on_capex", "da_total", "fs_revenue", "fs_ebit", "auto_cfo", "auto_cfi_operating", "auto_net_cash_flow", "auto_net_liquidity",
        "deliveries", "vehicle_sales", "asp_published_eur_k", "bev_share_published", "rd_cap_rate_published", "auto_ebitda_margin_published"]
for row in psql(f"select period_code, {', '.join(base)} from mart.v_panel_hy order by sort_key"):
    for c, v in zip(base, row[1:]):
        rec("R5 export integrity", f"FactHalfYear.{c} {row[0]}", f(FH[row[0]][c]), f(v), 0)

# ------------------------------------------------------------------ report
out = []
groups = sorted({g for g, *_ in RES})
out.append(f"{'check group':32} {'compared':>8} {'failed':>6} {'max |diff|':>12}")
for g in groups:
    rs = [r for r in RES if r[0] == g]; fails = [r for r in rs if not r[6]]
    mx = max([r[4] for r in rs if r[4] is not None] or [0])
    out.append(f"{g:32} {len(rs):>8} {len(fails):>6} {mx:>12.2e}")
tot_f = sum(1 for r in RES if not r[6])
out.append(f"\nTOTAL comparisons: {len(RES)}   failures: {tot_f}   RESULT: {'PASS' if tot_f == 0 else 'FAIL'}")
if XL_NOTE:
    out.append(f"\nFINDING (Excel presentation, not a Power BI difference): {len(XL_NOTE)} Excel cells show a value where no prior-year half exists;")
    out.append("Excel's own audit checks these rows only for comparable periods (F:J). Power BI returns BLANK. Excluded from R3:")
    out += ["  " + n for n in XL_NOTE]
for r in [r for r in RES if not r[6]][:40]:
    out.append(f"  FAIL {r[0]} | {r[1]} | replica={r[2]} reference={r[3]} tol={r[5]}")

# key values the dashboard will display (latest period) - for the spec and narrative
L = LATEST; Q = PRIOR[L]; e = rec_eff(L)
KEY = {
 "period": L, "prior": Q, "revenue": R.Revenue(L), "L0": R.L0(L), "L1": R.L1(L), "L2": R.L2(L), "L3": R.L3(L),
 "M0": R.M0(L), "M1": R.M1(L), "M2": R.M2(L), "M3": R.M3(L), "M0_py": R.M0(Q), "M1_py": R.M1(Q), "M2_py": R.M2(Q),
 "yoy_L0": YOY(R.L0, L), "yoy_L1": YOY(R.L1, L), "yoy_M0_pp": YOY(R.M0, L, "pp"), "yoy_M1_pp": YOY(R.M1, L, "pp"), "yoy_M2_pp": YOY(R.M2, L, "pp"),
 "bridge": {s: bridge(L, s) for s in ["S0_START", "S1_EXC", "S2_TAR", "S3_CAPDEV", "S4_DA", "S5_FS", "S6_RES"]},
 "bridge_pp": {s: bridge(L, s, "pp") for s in ["S1_EXC", "S2_TAR", "S3_CAPDEV", "S4_DA", "S5_FS", "S6_RES", "S7_DENOM"]},
 "res_alt_a04": sens(L, "SENS_A04_RES"), "L3_alt_a04": sens(L, "SENS_A04_L3"),
 "one_offs_share": div(bridge(L, "S1_EXC") + bridge(L, "S2_TAR"), YOY(R.L0, L)),
 "rd_total": R.RDT(L), "cap_dev": R.CapDev(L), "rd_pl": R.RDPL(L), "cap_rate": R.CapRate(L), "cap_rate_py": R.CapRate(Q), "cap_fx_sql": f(FH[L]["sql_ebit_effect_of_cap_rate_change"]),
 "net_cap": R.NetCap(L), "clean_da": R.CleanDA(L), "da_total": s0("da_total", L), "impairment": s0("impairment_used", L),
 "auto_rev": R.AutoRev(L), "vs": R.VS(L), "vs_yoy": YOY(R.VS, L, "pct"), "deliv": R.Deliv(L), "asp": R.ASP(L), "asp_yoy": YOY(R.ASP, L, "pct"),
 "vol_fx": f(FH[L]["sql_volume_effect"]), "asp_fx": f(FH[L]["sql_asp_effect"]),
 "ncf": R.NCF(L), "ncf_m": R.NCFM(L), "ncf_m_py": R.NCFM(Q), "ebitda_m": R.EBM(L), "cfo_ebitda": R.CFOE(L), "proxy": R.Proxy(L),
 "g": {lv: (g_min(f"h2_margin_{lv}_pct"), g_mid(f"h2_margin_{lv}_pct"), g_max(f"h2_margin_{lv}_pct")) for lv in ("l0", "l1", "l2")},
 "prior_h2_M0": R.M0("2025-H2"), "prior_h2_M1": R.M1("2025-H2"),
 "ga": {w: (min(f(r[c]) for r in GA.values()) / 100, f(GA["mid/mid"][c]) / 100, max(f(r[c]) for r in GA.values()) / 100)
        for w, c in (("ebitda", "h2_auto_ebitda_margin_pct"), ("ncf", "h2_auto_ncf_margin_pct"))},
 "mix": {dim: {r["member_name"]: (mix_share(L, dim, r["member_id"]), mix_share(Q, dim, r["member_id"])) for r in MIX if r["period_code"] == L and r["dimension"] == dim} for dim in ("model", "region")},
 "trend": {p: {"M0": R.M0(p), "M1": R.M1(p), "M2": R.M2(p), "M3": R.M3(p), "rev": R.Revenue(p), "L0": R.L0(p), "L1": R.L1(p), "L2": R.L2(p),
               "vs": R.VS(p), "asp": R.ASP(p), "ncf_m": R.NCFM(p), "rd_t": R.RDT(p), "cap": R.CapDev(p), "rd_pl": R.RDPL(p), "cap_rate": R.CapRate(p),
               "clean_da": R.CleanDA(p), "imp": s0("impairment_used", p), "ncf": R.NCF(p), "proxy": R.Proxy(p), "ebitda_m": R.EBM(p),
               "vol_fx": f(FH[p]["sql_volume_effect"]), "asp_fx": f(FH[p]["sql_asp_effect"]), "cap_fx": f(FH[p]["sql_ebit_effect_of_cap_rate_change"])} for p in P},
 "grid": {sid: {"l0": f(r["h2_margin_l0_pct"]), "l1": f(r["h2_margin_l1_pct"]), "l2": f(r["h2_margin_l2_pct"])} for sid, r in GD.items()},
 "n_ref": len(REF), "n_ref_pass": sum(1 for q in QROWS if q[5] == "PASS"), "auto_rev_py": R.AutoRev(Q),
 "qa_sample": [q for q in QROWS if (q[0], q[1]) in {("ebit_l1_ex_realignment", L), ("margin_l1_pct", L), ("eff_residual_other_drivers", L),
                                                    ("asp_calc_eur_k", L), ("g_h2_margin_l1_pct", "mid/mid/mid"), ("gs_l0_min", "ALL")}],
}
json.dump(KEY, open(os.path.join(ROOT, "model", "key_values.json"), "w"), indent=1)
open(os.path.join(HERE, "PCL_06_reconciliation_report.txt"), "w").write("\n".join(out) + "\n")
print("\n".join(out))
