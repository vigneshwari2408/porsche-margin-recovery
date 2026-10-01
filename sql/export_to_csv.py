"""
PCL Phase 3 - export the Phase 1/2 workbooks to flat CSV files for the SQL staging layer.

Inputs  (same folder or give paths):
  PCL_01_Sources.xlsx        -> source_register.csv, variable_dictionary.csv, audit_findings.csv
  PCL_02_Manual_Inputs.xlsx  -> manual_inputs.csv, data_gaps.csv
  fact_sheet_long.csv        -> copied unchanged (produced by pcl_extract_factsheets.py)
Output folder: data/

Rules:
- Nothing is recalculated except value_std (EUR bn -> EUR m, the same rule as the workbook formula).
- Blank inputs are NOT written to manual_inputs.csv; they go to data_gaps.csv with their written reason.
- Periods are converted from 'H1_2023' / 'FY_2023' to the SQL period codes '2023-H1' / '2023-FY'.
"""
import csv, os, shutil, sys
import openpyxl

SRC_WB, INP_WB, FS_CSV, OUT = "PCL_01_Sources.xlsx", "PCL_02_Manual_Inputs.xlsx", "fact_sheet_long.csv", "data"

# Variable dictionary: classification used by SQL (aggregation rule decides what may be summed or derived)
CLASS = {  # var_id: (var_group, measure_type, aggregation, unit_std, is_memo)
 **{v: ("Group P&L", "flow", "SUM", "EUR m", False) for v in ["G01","G02","G03","G04","G05","G06","G07"]},
 "G08": ("Group P&L", "ratio", "NONE", "%", False),
 "G09": ("Group P&L", "per_share", "NONE", "EUR/share", False),
 "G10": ("Group P&L", "per_share", "NONE", "EUR/share", False),
 **{v: ("EBIT bridge", "flow", "SUM", "EUR m", False) for v in ["B01","B02","B03","B04"]},
 "B05": ("EBIT bridge", "flow", "SUM", "EUR m", True),
 "X01": ("Exceptional items", "flow", "SUM", "EUR m", False),
 "X02": ("Exceptional items", "flow", "SUM", "EUR m", True),
 "X03": ("Exceptional items", "flow", "SUM", "EUR m", False),
 "X04": ("Exceptional items", "flow", "SUM", "EUR m", False),
 "X05": ("Exceptional items", "flow", "SUM", "EUR m", True),
 **{v: ("Automotive segment", "flow", "SUM", "EUR m", False) for v in ["A01","A02","A03"]},
 "A04": ("Automotive segment", "ratio", "NONE", "%", False),
 **{v: ("R&D, capex, D&A", "flow", "SUM", "EUR m", False) for v in ["A05","A06","A08","A09","A10","A11","A12"]},
 "A07": ("R&D, capex, D&A", "ratio", "NONE", "%", False),
 "A13": ("Automotive cash", "flow", "SUM", "EUR m", False),
 "A13a": ("Automotive cash", "flow", "SUM", "EUR m", False),
 "A13b": ("Automotive cash", "flow", "SUM", "EUR m", False),
 "A14": ("Automotive cash", "stock", "LAST", "EUR m", False),
 "F01": ("Financial Services", "flow", "SUM", "EUR m", False),
 "F02": ("Financial Services", "flow", "SUM", "EUR m", False),
 "O01": ("Volumes", "flow", "SUM", "units", False),
 "O02": ("Volumes", "flow", "SUM", "units", False),
 "O03": ("Volumes", "flow", "SUM", "units", False),
 "O04": ("Volumes", "flow", "SUM", "units", False),
 "O05": ("Volumes", "flow", "SUM", "units", False),
 "O06": ("Unit economics", "ratio", "NONE", "EUR k", False),
 "O07": ("Volumes", "ratio", "NONE", "%", False),
 "O08": ("Volumes", "flow", "SUM", "units", False),
 "O09": ("Unit economics", "ratio", "NONE", "EUR k", False),
}
EXTRA = {  # fact-sheet variables not in the Phase 1 dictionary
 "G09": ("Basic/diluted EPS, ordinary share", "Group", "Earnings per ordinary share (fact sheet)."),
 "G10": ("Basic/diluted EPS, preferred share", "Group", "Earnings per preferred share (fact sheet)."),
 "A13a": ("Automotive cash flows from operating activities", "Auto", "Operating cash flow of the Automotive segment (fact sheet)."),
 "A13b": ("Automotive investing activities of current operations", "Auto", "Investing cash flow attributable to operating activities, Automotive (fact sheet)."),
}

def conv_period(p):  # 'H1_2023' -> '2023-H1'
    k, y = p.split("_"); return f"{y}-{k}"

def write(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh); w.writerow(header); w.writerows(rows)
    print(f"{path}: {len(rows)} rows")

def main(base="."):
    os.makedirs(os.path.join(base, OUT), exist_ok=True)
    wb = openpyxl.load_workbook(os.path.join(base, SRC_WB), data_only=True)
    # source register
    ws = wb["Source_Register"]; rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0]: rows.append([r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[10]])
    write(os.path.join(base, OUT, "source_register.csv"),
          ["source_id","tier","publisher","document_title","doc_type","reporting_period","publication_date","url","file_name","notes"], rows)
    # variable dictionary
    ws = wb["Definitions"]; rows = []; seen = set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]: continue
        g, mt, agg, u, memo = CLASS[r[0]]
        rows.append([r[0], r[1].strip(), g, r[3], mt, agg, u, memo, r[5], r[7] or ""]); seen.add(r[0])
    for v, (name, seg, d) in EXTRA.items():
        g, mt, agg, u, memo = CLASS[v]
        rows.append([v, name, g, seg, mt, agg, u, memo, d, ""])
    write(os.path.join(base, OUT, "variable_dictionary.csv"),
          ["var_id","var_name","var_group","segment","measure_type","aggregation","unit_std","is_memo","definition","comparability_note"], rows)
    # audit findings
    ws = wb["Audit_Findings"]
    rows = [list(r[:6]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    write(os.path.join(base, OUT, "audit_findings.csv"), ["af_id","question","evidence","finding","status","consequence"], rows)

    # manual inputs
    wb2 = openpyxl.load_workbook(os.path.join(base, INP_WB), data_only=True)
    ws = wb2["Inputs"]; rows, gaps = [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        rec, var, name, per, src, page, fb, unit, val, _vstd, _su, label, rnd, who, dc, note = r[:16]
        basis = r[18] if len(r) > 18 else None
        if not rec: continue
        if val is None:
            gaps.append([var, conv_period(per), src, page, note]); continue
        vstd = val * 1000 if unit == "EUR bn" else val
        sunit = "EUR m" if unit == "EUR bn" else unit
        rows.append([rec, var, conv_period(per), src, page, unit, val, vstd, sunit, label, rnd or "", dc or "", note or "", basis or ""])
    write(os.path.join(base, OUT, "manual_inputs.csv"),
          ["rec_id","var_id","period_code","source_id","source_page","unit_as_shown","value_as_shown","value_std","std_unit","value_label","rounding","double_checked","note","verification_basis"], rows)
    write(os.path.join(base, OUT, "data_gaps.csv"), ["var_id","period_code","source_id","source_page","reason"], gaps)
    shutil.copy(os.path.join(base, FS_CSV), os.path.join(base, OUT, "fact_sheet_long.csv"))
    print("fact_sheet_long.csv copied")

if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else ".")
