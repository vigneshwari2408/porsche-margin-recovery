"""
PCL Phase 5 - build PCL_05_Porsche_Margin_Model.xlsx from the validated SQL outputs.

Design
- Controlled imports (blue values): Data_Long (core.fact_financial, analysis window), Data_Mix (deliveries),
  Guidance inputs, published reference ratios, SQL_Outputs (mart views, reconciliation only).
- Every analytical number is an Excel formula built from the imports (black) or a cross-sheet link (green).
- Audit_Checks re-computes identities and reconciles every key Excel result to the SQL mart views.
- Data_Panel uses two-criteria SUMIFS/COUNTIFS on Data_Long period_code + var_id (no helper key, no INDEX/MATCH).
"""
import os
import csv, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import LineChart, BarChart, Reference

SRC = sys.argv[1]            # folder with exported CSVs
OUT = sys.argv[2]            # output xlsx

def rd(name):
    with open(f"{SRC}/{name}", encoding="utf-8") as fh:
        return list(csv.reader(fh))

# ---------------------------------------------------------------- styles
AR = "Arial"
F_N = Font(name=AR, size=10)
F_B = Font(name=AR, size=10, bold=True)
F_IN = Font(name=AR, size=10, color="0000FF")                 # hard-coded input / import
F_LINK = Font(name=AR, size=10, color="008000")               # cross-sheet link
F_TITLE = Font(name=AR, size=14, bold=True, color="1F3864")
F_SUB = Font(name=AR, size=10, italic=True, color="595959")
F_H = Font(name=AR, size=10, bold=True, color="FFFFFF")
F_SEC = Font(name=AR, size=10, bold=True, color="1F3864")
F_WARN = Font(name=AR, size=10, bold=True, color="C00000")
FILL_H = PatternFill("solid", start_color="1F3864")
FILL_SEC = PatternFill("solid", start_color="D9E1F2")
FILL_SUPP = PatternFill("solid", start_color="EDEDED")
FILL_KEY = PatternFill("solid", start_color="FFF2CC")
FILL_SCEN = PatternFill("solid", start_color="FCE4D6")
FILL_DER = PatternFill("solid", start_color="FBE5D6")
thin = Side(style="thin", color="BFBFBF")
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)

FMT = {
    "eur":  '#,##0;(#,##0);"-"',
    "eur1": '#,##0.0;(#,##0.0);"-"',
    "pct":  '0.0%;(0.0%);"-"',
    "pct2": '0.00%;(0.00%);"-"',
    "pp":   '+0.00" pp";-0.00" pp";0.00" pp"',
    "k":    '#,##0.0;(#,##0.0);"-"',
    "units": '#,##0;(#,##0);"-"',
    "x":    '0.00"x";(0.00"x");"-"',
    "raw":  'General',
    "txt":  '@',
}

P = ["2023-H1", "2023-H2", "2024-H1", "2024-H2", "2025-H1", "2025-H2", "2026-H1"]
PC = ["D", "E", "F", "G", "H", "I", "J"]        # period columns in every analysis sheet
PRIOR = {PC[i]: (PC[i - 2] if i >= 2 else None) for i in range(7)}
HDR_ROW = 5                                      # period header row in analysis sheets
FIRST = 6                                        # first data row in analysis sheets

wb = Workbook()
wb.remove(wb.active)
ORDER = ["README", "Data_Panel", "Data_Long", "Data_Mix", "SQL_Outputs", "KPI_Overview", "EBIT_Analysis",
         "RD_DA", "Commercial", "Cash_Crosscheck", "Guidance_Scenarios", "Audit_Checks", "Audit_Register"]
WS = {n: wb.create_sheet(n) for n in ORDER}

def title(ws, t, sub, widths):
    ws["A1"] = t; ws["A1"].font = F_TITLE
    ws["A2"] = sub; ws["A2"].font = F_SUB
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False

def hdr(ws, row, values, start_col=1, fill=FILL_H, font=F_H):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start_col + i, value=v)
        c.font = font; c.fill = fill; c.alignment = CEN; c.border = BD

SCEN_LABEL = "Scenario combination - not a forecast; no probability"
L3_HEADING = "E. SUPPLEMENTARY HYBRID SENSITIVITY (L3) - not underlying group profitability; NOT a headline KPI"
# ================================================================ Data_Long
ws = WS["Data_Long"]
title(ws, "Data_Long - controlled import of core.fact_financial (analysis window)",
      "Source: PostgreSQL porsche_lens, core.fact_financial filtered on dim_period.in_analysis_window (2023-H1..2026-H1). "
      "Blue = imported value (do not edit). K = rows with the same period_code + var_id (must be 1); L = value read back from Data_Panel; M = |L - D| (must be 0).",
      {"A": 10, "B": 8, "C": 44, "D": 16, "E": 9, "F": 9, "G": 12, "H": 9, "I": 70, "J": 36, "K": 12, "L": 16, "M": 12})
rows = rd("fact_long.csv")[1:]
hdr(ws, 5, ["period_code", "var_id", "var_name", "value", "unit_std", "value_label", "origin", "source_id",
            "source_ref", "rounding", "dup_count", "panel_readback", "readback_diff"])
DL_FIRST, DL_LAST = 6, 5 + len(rows)
for i, r in enumerate(rows):
    rr = DL_FIRST + i
    for j, v in enumerate(r):
        c = ws.cell(row=rr, column=j + 1, value=(float(v) if j == 3 else v))
        c.font = F_IN if j == 3 else F_N
        if j == 3: c.number_format = '#,##0.000000'
    ws.cell(row=rr, column=11, value=f'=COUNTIFS($A${DL_FIRST}:$A${DL_LAST},A{rr},$B${DL_FIRST}:$B${DL_LAST},B{rr})').font = F_N
ws.freeze_panes = "A6"
ws.auto_filter.ref = f"A5:M{DL_LAST}"
ws["A3"] = f"Rows imported: {len(rows)}   |   Export query: see 03_SQL/12_export_mart.sql conventions; re-import by replacing rows {DL_FIRST}-{DL_LAST}."
ws["A3"].font = F_SUB
DLR = lambda col: f"Data_Long!${col}${DL_FIRST}:${col}${DL_LAST}"

# ================================================================ Data_Mix
ws = WS["Data_Mix"]
title(ws, "Data_Mix - controlled import of deliveries by model line and region",
      "Source: core.fact_deliveries_model / core.fact_deliveries_region (fact sheet S025; H2 = FY - H1, DERIVED). Blue = imported.",
      {"A": 10, "B": 10, "C": 10, "D": 22, "E": 10, "F": 12, "G": 10, "H": 9, "I": 40})
mrows = rd("mix.csv")[1:]
hdr(ws, 5, ["period_code", "dimension", "member_id", "member_name", "top_level", "deliveries", "value_label", "source_id", "source_ref"])
MX_FIRST, MX_LAST = 6, 5 + len(mrows)
for i, r in enumerate(mrows):
    rr = MX_FIRST + i
    for j, v in enumerate(r):
        val = int(v) if j == 5 else ("TRUE" if (j == 4 and v == "t") else ("FALSE" if j == 4 else v))
        c = ws.cell(row=rr, column=j + 1, value=val)
        c.font = F_IN if j == 5 else F_N
        if j == 5: c.number_format = FMT["units"]
ws.freeze_panes = "A6"; ws.auto_filter.ref = f"A5:I{MX_LAST}"
MXR = lambda col: f"Data_Mix!${col}${MX_FIRST}:${col}${MX_LAST}"

# ================================================================ Data_Panel
PANEL_VARS = [
 ("G01","Group sales revenue","EUR m"),("G02","Cost of sales","EUR m"),("G03","Gross profit","EUR m"),
 ("G04","Distribution expenses","EUR m"),("G05","Administrative expenses","EUR m"),("G06","Net other operating result","EUR m"),
 ("G07","Operating profit (EBIT), group","EUR m"),
 ("A01","Automotive sales revenue","EUR m"),("A02","Automotive operating profit","EUR m"),("A03","Automotive EBITDA","EUR m"),
 ("A04","Automotive EBITDA margin (published)","%"),
 ("A05","Automotive R&D costs (total)","EUR m"),("A06","Capitalised development costs","EUR m"),("A07","Capitalisation rate (published)","%"),
 ("A08","Expensed R&D","EUR m"),("A09","Amortisation of capitalised R&D","EUR m"),("A10","Automotive capex","EUR m"),
 ("A11","Depreciation on capex (incl. impairments)","EUR m"),("A12","Automotive D&A total (incl. impairments)","EUR m"),
 ("A13a","Automotive cash flow from operating activities","EUR m"),("A13b","Automotive investing of current operations","EUR m"),
 ("A13","Automotive net cash flow","EUR m"),("A14","Automotive net liquidity (period end)","EUR m"),
 ("F01","Financial Services sales revenue","EUR m"),("F02","Financial Services operating profit","EUR m"),
 ("X01","Strategic realignment incl. battery (net charge)","EUR m"),("X02","  of which battery activities (memo)","EUR m"),
 ("X03","US import tariffs (expense)","EUR m"),("X04","Provision release (benefit)","EUR m"),("X05","Impairments inside automotive D&A","EUR m"),
 ("B01","Porsche bridge: gross margin w/o R&D (YoY)","EUR m"),("B02","Porsche bridge: R&D (YoY)","EUR m"),
 ("B03","Porsche bridge: SG&A (YoY)","EUR m"),("B04","Porsche bridge: Other (YoY)","EUR m"),("B05","Porsche bridge: extraordinary (memo)","EUR m"),
 ("O01","Deliveries (retail)","units"),("O04","Vehicle sales (wholesale)","units"),("O06","ASP per vehicle sold (published)","EUR k"),
 ("O07","BEV share (published)","%"),
]
ws = WS["Data_Panel"]
title(ws, "Data_Panel - half-year panel built by formula from Data_Long",
      "Each cell = two-criteria lookup on Data_Long: SUMIFS(value, period_code = column header, var_id = row id), guarded by COUNTIFS (0 matches -> blank; Data_Long!K proves each pair is unique). Blank = not published. "
      "Shaded cells are DERIVED (H2 = FY - H1 or identity); label grid below states FACT / DERIVED per cell.",
      {"A": 8, "B": 46, "C": 8, **{c: 12 for c in PC}, "K": 44})
ws["A3"] = "Money in EUR m; % variables as published (e.g. 18.3 = 18.3%). Analysis sheets convert."; ws["A3"].font = F_SUB
hdr(ws, HDR_ROW, ["var_id", "variable", "unit"] + P + ["sources used (from Data_Long)"])
PROW = {}
srcs = {}
for r in rows:
    srcs.setdefault(r[1], set()).add(r[7] if r[7] else "SQL derived")
for i, (v, name, unit) in enumerate(PANEL_VARS):
    rr = FIRST + i; PROW[v] = rr
    ws.cell(row=rr, column=1, value=v).font = F_B
    ws.cell(row=rr, column=2, value=name).font = F_N
    ws.cell(row=rr, column=3, value=unit).font = F_N
    for c in PC:
        crit = f'{DLR("A")},{c}${HDR_ROW},{DLR("B")},$A{rr}'
        f = f'=IF(COUNTIFS({crit})=0,"",SUMIFS({DLR("D")},{crit}))'
        cell = ws[f"{c}{rr}"]; cell.value = f; cell.font = F_N
        cell.number_format = FMT["units"] if unit == "units" else ('0.0' if unit in ("%", "EUR k") else FMT["eur"])
        cell.border = BD
    ws.cell(row=rr, column=11, value=", ".join(sorted(srcs.get(v, [])))).font = F_SUB
for j, p in enumerate(P):
    ws[f"{PC[j]}{HDR_ROW}"].value = p
PANEL_LAST = FIRST + len(PANEL_VARS) - 1
LBL_TOP = PANEL_LAST + 3
ws.cell(row=LBL_TOP - 1, column=1, value="Label grid: FACT = published by Porsche; DERIVED = calculated in SQL from published values (see Data_Long source_ref).").font = F_SEC
hdr(ws, LBL_TOP, ["var_id", "variable", ""] + P)
LROW = {}
for i, (v, name, unit) in enumerate(PANEL_VARS):
    rr = LBL_TOP + 1 + i; LROW[v] = rr
    ws.cell(row=rr, column=1, value=v).font = F_N
    ws.cell(row=rr, column=2, value=name).font = F_N
    for c in PC:
        crit = f'{DLR("A")},{c}${HDR_ROW},{DLR("B")},$A{rr}'
        ws[f"{c}{rr}"] = f'=IF(COUNTIFS({crit})=0,"not published",IF(COUNTIFS({crit},{DLR("F")},"DERIVED")>0,"DERIVED","FACT"))'
        ws[f"{c}{rr}"].font = F_N; ws[f"{c}{rr}"].alignment = CEN
    # shade derived cells in the value block
off = LBL_TOP + 1 - FIRST
ws.conditional_formatting.add(f"D{FIRST}:J{PANEL_LAST}",
    FormulaRule(formula=[f'D{FIRST + off}="DERIVED"'], fill=FILL_DER))
ws.freeze_panes = f"D{FIRST}"
# --- round-trip: every Data_Long row read back from the Data_Panel grid (2-D, period header x var_id)
_dl = WS["Data_Long"]
_PV = f"Data_Panel!$A${FIRST}:$A${PANEL_LAST}"; _PH = f"Data_Panel!$D${HDR_ROW}:$J${HDR_ROW}"; _PG = f"Data_Panel!$D${FIRST}:$J${PANEL_LAST}"
for rr in range(DL_FIRST, DL_LAST + 1):
    _rb = '"period not in panel"'
    for _c in reversed(PC):
        _rb = f'IF(A{rr}=Data_Panel!${_c}${HDR_ROW},SUMIFS(Data_Panel!${_c}${FIRST}:${_c}${PANEL_LAST},{_PV},B{rr}),{_rb})'
    _dl.cell(row=rr, column=12, value=f'=IF(COUNTIF({_PV},B{rr})=0,"not in panel",{_rb})').font = F_N
    _dl.cell(row=rr, column=13, value=f'=IF(L{rr}="not in panel",0,ABS(N(L{rr})-D{rr}))').font = F_N
    _dl.cell(row=rr, column=12).number_format = '#,##0.000000'; _dl.cell(row=rr, column=13).number_format = '0.000000'

def pref(v, col):     # reference to Data_Panel
    return f"Data_Panel!{col}{PROW[v]}"

# ================================================================ sheet row DSL
# tokens: {c} current col, {p} prior col; [KEY] / [KEY@p] same sheet; [Sheet:KEY] / [Sheet:KEY@p]; <VAR> / <VAR@p> Data_Panel
REG = {}   # (sheet, key) -> row

def layout(sheet, spec, start=FIRST):
    r = start
    for item in spec:
        if item.get("key"):
            REG[(sheet, item["key"])] = r
        item["_row"] = r
        r += 1 + item.get("gap", 0)
    return r

def render(expr, sheet, col):
    pcol = PRIOR[col]
    need_p = "@p" in expr or "{p}" in expr
    if need_p and pcol is None:
        return None
    def sub_ref(m):
        sh, key, at = m.group(1), m.group(2), m.group(3)
        c = pcol if at else col
        target = sh[:-1] if sh else sheet
        row = REG[(target, key)]
        return (f"{target}!{c}{row}" if sh else f"{c}{row}")
    def sub_panel(m):
        v, at = m.group(1), m.group(2)
        return pref(v, pcol if at else col)
    e = re.sub(r"\[([A-Za-z_]+:)?([A-Za-z0-9_]+)(@p)?\]", sub_ref, expr)
    e = re.sub(r"<([A-Z0-9a-z]+)(@p)?>", sub_panel, e)
    e = e.replace("{c}", col).replace("{p}", pcol or "")
    return e

def write_sheet(sheet, spec, label_w=50):
    ws = WS[sheet]
    for item in spec:
        r = item["_row"]
        kind = item.get("kind", "row")
        if kind == "sec":
            c = ws.cell(row=r, column=1, value=item["label"]); c.font = item.get("font", F_SEC)
            for col in range(1, 12):
                ws.cell(row=r, column=col).fill = item.get("fill", FILL_SEC)
            continue
        if kind == "note":
            c = ws.cell(row=r, column=1, value=item["label"]); c.font = item.get("font", F_SUB)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=11)
            c.alignment = WR
            ws.row_dimensions[r].height = item.get("h", 30)
            continue
        ws.cell(row=r, column=1, value=item.get("id", "")).font = F_N
        ws.cell(row=r, column=2, value=item["label"]).font = F_B if item.get("bold") else F_N
        ws.cell(row=r, column=3, value=item.get("unit", "")).font = F_N
        only = item.get("only")
        for j, col in enumerate(PC):
            if only and P[j] not in only:
                continue
            if "values" in item:            # hard-coded inputs per period
                v = item["values"].get(P[j])
                if v is None: continue
                cell = ws[f"{col}{r}"]; cell.value = v; cell.font = F_IN
            else:
                e = render(item["f"], sheet, col)
                if e is None: continue
                cell = ws[f"{col}{r}"]; cell.value = "=" + e
                is_link = bool(re.fullmatch(r"=?[A-Za-z_]+![A-Z]+\d+", e))
                cell.font = F_LINK if is_link else F_N
            cell.number_format = FMT[item.get("fmt", "eur")]
            cell.border = BD
            if item.get("fill"): cell.fill = item["fill"]
            if item.get("bold"): cell.font = Font(name=AR, size=10, bold=True, color=cell.font.color)
        if item.get("src"):
            c = ws.cell(row=r, column=11, value=item["src"]); c.font = F_SUB; c.alignment = WR
        if item.get("fill"):
            for col in ("A", "B", "C"):
                ws[f"{col}{r}"].fill = item["fill"]

def sheet_frame(sheet, t, sub, question):
    ws = WS[sheet]
    title(ws, t, sub, {"A": 7, "B": 52, "C": 8, **{c: 12 for c in PC}, "K": 70})
    ws["A3"] = "Business question: " + question; ws["A3"].font = F_B
    hdr(ws, HDR_ROW, ["ID", "Measure", "Unit"] + [None] * 7 + ["Formula / source lineage"])
    for j, c in enumerate(PC):
        ws[f"{c}{HDR_ROW}"] = f"=Data_Panel!{c}${HDR_ROW}"
        ws[f"{c}{HDR_ROW}"].font = Font(name=AR, size=10, bold=True, color="FFFFFF")
    ws.freeze_panes = f"D{FIRST}"

# ================================================================ EBIT_Analysis spec
PUB_ROS = {"2024-H1": 0.157, "2025-H1": 0.055, "2026-H1": 0.078}
EB = [
 {"kind": "sec", "label": "A. Reported group profitability"},
 {"key": "rev", "id": "E01", "label": "Group sales revenue", "unit": "EUR m", "f": "<G01>", "src": "Data_Panel G01 <- S025 fact sheet"},
 {"key": "l0", "id": "E02", "label": "L0 Reported EBIT", "unit": "EUR m", "f": "<G07>", "bold": True, "src": "Data_Panel G07 <- S025 fact sheet"},
 {"key": "m0", "id": "E03", "label": "L0 Reported margin (return on sales)", "unit": "%", "fmt": "pct", "f": "[l0]/[rev]", "bold": True, "src": "E02 / E01"},
 {"key": "ros_pub", "id": "E04", "label": "Published group return on sales (reference)", "unit": "%", "fmt": "pct", "values": PUB_ROS,
  "src": "Input: S009 slide 27 (H1 2024), S016 slide 5 (H1 2025, H1 2026). Used only by check C28."},
 {"key": "ros_diff", "id": "E05", "label": "Check: calculated - published margin", "unit": "pp", "fmt": "pp",
  "f": 'IF([ros_pub]="",0,([m0]-[ros_pub])*100)', "src": "0 where no published value"},
 {"kind": "sec", "label": "B. Exceptional items disclosed by Porsche (EUR m, positive = expense)"},
 {"key": "x01", "id": "E06", "label": "Strategic realignment incl. battery (net), used", "unit": "EUR m", "f": 'IF(<X01>="",0,<X01>)',
  "src": "Data_Panel X01 (S016 p.6; S011 p.6 FY2025 = 2.4 + 0.7, DERIVED). Blank -> 0 (assumption A-01)."},
 {"key": "x01f", "id": "", "label": "   status", "unit": "", "fmt": "txt", "f": 'IF(<X01>="","A-01: none disclosed -> 0","disclosed")'},
 {"key": "x03", "id": "E07", "label": "US import tariffs, used", "unit": "EUR m", "f": 'IF(<X03>="",0,<X03>)',
  "src": "Data_Panel X03 (S010 p.5; S011 p.6; S016 p.6). Blank -> 0 (assumption A-02)."},
 {"key": "x03f", "id": "", "label": "   status", "unit": "", "fmt": "txt", "f": 'IF(<X03>="","A-02: none before Apr 2025 -> 0","disclosed")'},
 {"kind": "note", "label": "Impairments (Data_Panel X05) are NOT added here: they are already inside the realignment/battery line. Adding them would double count."},
 {"kind": "sec", "label": "C. Adjusted EBIT - HEADLINE LEVELS"},
 {"key": "l1", "id": "E08", "label": "L1 EBIT excl. realignment & battery", "unit": "EUR m", "f": "[l0]+[x01]", "bold": True, "fill": FILL_KEY, "src": "E02 + E06"},
 {"key": "m1", "id": "E09", "label": "L1 margin", "unit": "%", "fmt": "pct", "f": "[l1]/[rev]", "bold": True, "fill": FILL_KEY, "src": "E08 / E01"},
 {"key": "l2", "id": "E10", "label": "L2 EBIT excl. realignment, battery & US tariffs", "unit": "EUR m", "f": "[l1]+[x03]", "bold": True, "fill": FILL_KEY, "src": "E08 + E07"},
 {"key": "m2", "id": "E11", "label": "L2 margin", "unit": "%", "fmt": "pct", "f": "[l2]/[rev]", "bold": True, "fill": FILL_KEY, "src": "E10 / E01"},
 {"key": "band", "id": "E12", "label": "Rounding band of adjustments (+/-)", "unit": "EUR m", "f": '(IF(<X01>="",0,50)+IF(<X03>="",0,50))*IF(RIGHT({c}$5,2)="H2",2,1)',
  "src": "One-offs printed to EUR 0.1bn -> +/-50m each; doubled for derived H2 (A-09)"},
 {"kind": "sec", "label": "D. Year-on-year change (same half of prior year)"},
 {"key": "d_l0", "id": "E13", "label": "Change in L0 EBIT", "unit": "EUR m", "f": "[l0]-[l0@p]"},
 {"key": "d_m0", "id": "E14", "label": "Change in L0 margin", "unit": "pp", "fmt": "pp", "f": "([m0]-[m0@p])*100"},
 {"key": "d_l1", "id": "E15", "label": "Change in L1 EBIT", "unit": "EUR m", "f": "[l1]-[l1@p]", "bold": True},
 {"key": "d_m1", "id": "E16", "label": "Change in L1 margin", "unit": "pp", "fmt": "pp", "f": "([m1]-[m1@p])*100", "bold": True},
 {"key": "d_l2", "id": "E17", "label": "Change in L2 EBIT", "unit": "EUR m", "f": "[l2]-[l2@p]"},
 {"key": "d_m2", "id": "E18", "label": "Change in L2 margin", "unit": "pp", "fmt": "pp", "f": "([m2]-[m2@p])*100"},
 {"kind": "sec", "label": L3_HEADING, "fill": FILL_SUPP, "font": F_WARN},
 {"kind": "note", "label": "L3 = L1 + clean automotive D&A - automotive capitalised development costs: an EBITDA-type figure in which total automotive R&D costs are expensed as incurred. "
                           "It mixes group EBIT with automotive-only items, so it is shown only as a sensitivity.", "h": 30},
 {"key": "daclean", "id": "E19", "label": "Clean automotive D&A (from RD_DA)", "unit": "EUR m", "f": "[RD_DA:daclean]", "fill": FILL_SUPP},
 {"key": "capdev", "id": "E20", "label": "Capitalised development costs", "unit": "EUR m", "f": "<A06>", "fill": FILL_SUPP},
 {"key": "l3", "id": "E21", "label": "L3 hybrid sensitivity", "unit": "EUR m", "f": "[l1]+[daclean]-[capdev]", "fill": FILL_SUPP, "src": "E08 + E19 - E20"},
 {"key": "m3", "id": "E22", "label": "L3 hybrid margin (sensitivity only)", "unit": "%", "fmt": "pct", "f": "[l3]/[rev]", "fill": FILL_SUPP},
 {"key": "l3alt", "id": "E23", "label": "L3 if H1 2026 impairments are not inside D&A (A-04 false)", "unit": "EUR m", "f": "[l3]+[RD_DA:impused]", "only": ["2026-H1"], "fill": FILL_SUPP},
 {"kind": "sec", "label": "F. Own year-on-year EBIT bridge (EUR m)"},
 {"key": "b_prior", "id": "E24", "label": "EBIT same half prior year", "unit": "EUR m", "f": "[l0@p]"},
 {"key": "b_d", "id": "E25", "label": "Change in reported EBIT", "unit": "EUR m", "f": "[l0]-[b_prior]", "bold": True},
 {"key": "e_exc", "id": "E26", "label": "  Exceptional items (realignment/battery, net)", "unit": "EUR m", "f": "-([x01]-[x01@p])", "src": "-(change in E06)"},
 {"key": "e_tar", "id": "E27", "label": "  US import tariffs", "unit": "EUR m", "f": "-([x03]-[x03@p])", "src": "-(change in E07)"},
 {"key": "e_rdc", "id": "E28", "label": "  Capitalised development costs", "unit": "EUR m", "f": "<A06>-<A06@p>", "src": "change in A06 (more capitalisation = less cost in P&L)"},
 {"key": "e_da", "id": "E29", "label": "  Clean automotive D&A", "unit": "EUR m", "f": "-([daclean]-[daclean@p])", "src": "-(change in clean D&A); impairments excluded (no double count)"},
 {"key": "e_fs", "id": "E30", "label": "  Financial Services EBIT", "unit": "EUR m", "f": "<F02>-<F02@p>", "src": "change in F02"},
 {"key": "e_res", "id": "E31", "label": "  Residual / other EBIT drivers (not separated)", "unit": "EUR m", "f": "[b_d]-[e_exc]-[e_tar]-[e_rdc]-[e_da]-[e_fs]", "bold": True,
  "src": "E25 - E26..E30. NOT operating performance: mixes price, mix, volume, total R&D cost changes, material/supplier costs, SG&A, other operating result, FX, consolidation, one-off rounding"},
 {"key": "e_res_alt", "id": "E32", "label": "  Residual if A-04 false (H1 2026 impairments not in D&A)", "unit": "EUR m",
  "f": "[b_d]-[e_exc]-[e_tar]-[e_rdc]-([e_da]-[RD_DA:impused])-[e_fs]", "only": ["2026-H1"]},
 {"kind": "sec", "label": "G. Same bridge in margin points (pp) - closes exactly"},
 {"key": "p_d", "id": "E33", "label": "Change in reported margin", "unit": "pp", "fmt": "pp", "f": "[d_m0]", "bold": True},
 {"key": "p_den", "id": "E34", "label": "  Revenue denominator effect", "unit": "pp", "fmt": "pp", "f": "100*[b_prior]*(1/[rev]-1/[rev@p])", "src": "prior EBIT x (1/rev_cur - 1/rev_prior)"},
 {"key": "p_exc", "id": "E35", "label": "  Exceptional items", "unit": "pp", "fmt": "pp", "f": "100*[e_exc]/[rev]"},
 {"key": "p_tar", "id": "E36", "label": "  US import tariffs", "unit": "pp", "fmt": "pp", "f": "100*[e_tar]/[rev]"},
 {"key": "p_rdc", "id": "E37", "label": "  Capitalised development costs", "unit": "pp", "fmt": "pp", "f": "100*[e_rdc]/[rev]"},
 {"key": "p_da", "id": "E38", "label": "  Clean automotive D&A", "unit": "pp", "fmt": "pp", "f": "100*[e_da]/[rev]"},
 {"key": "p_fs", "id": "E39", "label": "  Financial Services EBIT", "unit": "pp", "fmt": "pp", "f": "100*[e_fs]/[rev]"},
 {"key": "p_res", "id": "E40", "label": "  Residual / other EBIT drivers", "unit": "pp", "fmt": "pp", "f": "100*[e_res]/[rev]", "bold": True},
 {"key": "p_chk", "id": "E41", "label": "Check: components - total (must be 0)", "unit": "pp", "fmt": "pp",
  "f": "[p_den]+[p_exc]+[p_tar]+[p_rdc]+[p_da]+[p_fs]+[p_res]-[p_d]"},
 {"kind": "sec", "label": "H. Porsche's own bridge buckets (reconciliation of the total)"},
 {"key": "pb1", "id": "E42", "label": "  Gross margin without R&D", "unit": "EUR m", "f": 'IF(<B01>="","",<B01>)', "src": "Data_Panel B01 (decks S005/S007/S009/S011/S016; H2 = FY - H1)"},
 {"key": "pb2", "id": "E43", "label": "  R&D", "unit": "EUR m", "f": 'IF(<B02>="","",<B02>)'},
 {"key": "pb3", "id": "E44", "label": "  SG&A", "unit": "EUR m", "f": 'IF(<B03>="","",<B03>)'},
 {"key": "pb4", "id": "E45", "label": "  Other", "unit": "EUR m", "f": 'IF(<B04>="","",<B04>)'},
 {"key": "pbsum", "id": "E46", "label": "Sum of Porsche buckets", "unit": "EUR m", "f": 'IF(<B01>="","",SUM([pb1],[pb2],[pb3],[pb4]))'},
 {"key": "pbdiff", "id": "E47", "label": "Own change in EBIT - Porsche buckets", "unit": "EUR m", "f": 'IF([pbsum]="",0,[b_d]-[pbsum])'},
 {"key": "pbtol", "id": "E48", "label": "Tolerance (Porsche rounding)", "unit": "EUR m", "f": 'IF(AND(RIGHT({c}$5,2)="H2",VALUE(LEFT({c}$5,4))<=2024),250,50)',
  "src": "H2 2024 derived from FY2024 buckets printed to EUR 0.1bn"},
]
EB_END = layout("EBIT_Analysis", EB)

# ================================================================ RD_DA spec
RDS = [
 {"kind": "sec", "label": "A. Automotive R&D - total R&D costs (Porsche definition) vs R&D charged to the P&L"},
 {"kind": "note", "label": "Total R&D costs = research costs + non-capitalisable development costs + capitalisable development costs (Porsche definition). "
                           "This is an accounting cost total, NOT a cash-flow figure."},
 {"key": "rev", "id": "R01", "label": "Automotive sales revenue", "unit": "EUR m", "f": "<A01>"},
 {"key": "rdt", "id": "R02", "label": "Total R&D costs", "unit": "EUR m", "f": "<A05>", "bold": True, "src": "A05: S011 slide 34 (FY), S016 slide 25 (H1 2024-26), S002 p.29 (H1 2023)"},
 {"key": "cap", "id": "R03", "label": "  Capitalised development costs", "unit": "EUR m", "f": "<A06>"},
 {"key": "exp", "id": "R04", "label": "  Expensed R&D", "unit": "EUR m", "f": "<A08>"},
 {"key": "am", "id": "R05", "label": "Amortisation of capitalised R&D", "unit": "EUR m", "f": "<A09>"},
 {"key": "pl", "id": "R06", "label": "R&D charged to the P&L (expensed + amortisation)", "unit": "EUR m", "f": "[exp]+[am]", "bold": True},
 {"key": "cr", "id": "R07", "label": "Capitalisation rate (calculated)", "unit": "%", "fmt": "pct", "f": "[cap]/[rdt]", "bold": True},
 {"key": "crp", "id": "R08", "label": "Capitalisation rate (published)", "unit": "%", "fmt": "pct", "f": 'IF(<A07>="","",<A07>/100)'},
 {"key": "crdiff", "id": "R09", "label": "Check: calculated - published", "unit": "pp", "fmt": "pp", "f": 'IF([crp]="",0,([cr]-[crp])*100)'},
 {"key": "netcap", "id": "R10", "label": "Net capitalisation (capitalised - amortised)", "unit": "EUR m", "f": "[cap]-[am]", "src": ">0: P&L charge below total R&D costs"},
 {"key": "rdt_pct", "id": "R11", "label": "Total R&D costs / automotive revenue", "unit": "%", "fmt": "pct", "f": "[rdt]/[rev]"},
 {"key": "pl_pct", "id": "R12", "label": "R&D P&L charge / automotive revenue", "unit": "%", "fmt": "pct", "f": "[pl]/[rev]"},
 {"key": "cf", "id": "R13", "label": "P&L charge at prior-year capitalisation rate (illustrative)", "unit": "EUR m", "f": "[rdt]*(1-[cr@p])+[am]",
  "src": "Counterfactual: same total R&D costs and amortisation, prior-year same-half rate"},
 {"key": "cfe", "id": "R14", "label": "EBIT effect of capitalisation-rate change (+ = extra cost)", "unit": "EUR m", "f": "[pl]-[cf]", "bold": True},
 {"key": "d_pl", "id": "R15", "label": "Change in R&D P&L charge", "unit": "EUR m", "f": "[pl]-[pl@p]"},
 {"key": "d_rdt", "id": "R16", "label": "Change in total R&D costs", "unit": "EUR m", "f": "[rdt]-[rdt@p]"},
 {"key": "d_am", "id": "R17", "label": "Change in amortisation", "unit": "EUR m", "f": "[am]-[am@p]"},
 {"key": "chk_rd", "id": "R18", "label": "Check: total - capitalised - expensed (rounding)", "unit": "EUR m", "f": "[rdt]-[cap]-[exp]"},
 {"kind": "note", "label": "H2 2025 amortisation (FY - H1) most likely contains impairments of capitalised development costs; Porsche does not disclose the split."},
 {"kind": "sec", "label": "B. Automotive D&A and impairment treatment"},
 {"key": "da", "id": "R19", "label": "Automotive D&A total (incl. impairments)", "unit": "EUR m", "f": "<A12>", "src": "A12: EBITDA tables and S011 slide 34 / S016 slide 25"},
 {"key": "dep", "id": "R20", "label": "  Depreciation on capex (incl. impairments)", "unit": "EUR m", "f": "<A11>", "src": "H1 2023 = A12 - A09 (DERIVED)"},
 {"key": "am2", "id": "R21", "label": "  Amortisation of capitalised R&D", "unit": "EUR m", "f": "[am]"},
 {"key": "impdisc", "id": "R22", "label": "Impairments disclosed inside D&A", "unit": "EUR m", "f": 'IF(<X05>="","",<X05>)',
  "src": "X05: S002 p.38, S006 p.38, S010 p.32 (Cellforce 295), S011 slide 34 fn (~1,200 FY2025), S017 p.29 (43 + 18)"},
 {"key": "impused", "id": "R23", "label": "Impairments removed (used)", "unit": "EUR m", "f": 'IF(<X05>="",0,<X05>)', "bold": True},
 {"key": "impst", "id": "R24", "label": "   treatment", "unit": "", "fmt": "txt",
  "f": 'IF(<X05>="","A-03: not disclosed -> 0",IF({c}$5="2026-H1","A-04: 61 assumed inside D&A (upper bound)",IF(AND(RIGHT({c}$5,2)="H2",<X05>>0),"H2 = FY - H1; split not disclosed","FACT")))'},
 {"key": "daclean", "id": "R25", "label": "Clean automotive D&A (excl. impairments)", "unit": "EUR m", "f": "[da]-[impused]", "bold": True, "fill": FILL_KEY},
 {"key": "da_pct", "id": "R26", "label": "D&A total / automotive revenue", "unit": "%", "fmt": "pct", "f": "[da]/[rev]"},
 {"key": "dac_pct", "id": "R27", "label": "Clean D&A / automotive revenue", "unit": "%", "fmt": "pct", "f": "[daclean]/[rev]", "bold": True},
 {"key": "depclean", "id": "R28", "label": "Clean depreciation on capex (where impairment line is known)", "unit": "EUR m", "fmt": "eur",
  "f": 'IF([impused]=0,[dep],IF({c}$5="2025-H1",[dep]-[impused],"n/a"))', "src": "H1 2025: Cellforce impairment sits in PP&E depreciation (S010 p.32)"},
 {"key": "d_da", "id": "R29", "label": "Change in D&A total", "unit": "EUR m", "f": "[da]-[da@p]"},
 {"key": "d_dac", "id": "R30", "label": "Change in clean D&A", "unit": "EUR m", "f": "[daclean]-[daclean@p]", "bold": True},
 {"key": "chk_da", "id": "R31", "label": "Check: D&A - amortisation - depreciation (rounding)", "unit": "EUR m", "f": "[da]-[am]-[dep]"},
]
RD_END = layout("RD_DA", RDS)

# ================================================================ Commercial spec
MODELS = [("911", "911"), ("718", "718 Boxster/Cayman"), ("CAY", "Cayenne"), ("MAC", "Macan (total)"), ("PAN", "Panamera"), ("TAY", "Taycan")]
REGIONS = [("NA", "North America"), ("EU", "Europe excl. Germany"), ("DE", "Germany"), ("CN", "China"), ("OEM", "Overseas & Emerging Markets")]
CM = [
 {"kind": "sec", "label": "A. Revenue per vehicle sold (Porsche ASP definition) and volume"},
 {"key": "rev", "id": "V01", "label": "Automotive sales revenue", "unit": "EUR m", "f": "<A01>"},
 {"key": "vs", "id": "V02", "label": "Vehicle sales (wholesale units)", "unit": "units", "fmt": "units", "f": "<O04>", "src": "O04: S025 fact sheet"},
 {"key": "del", "id": "V03", "label": "Deliveries to customers (retail units)", "unit": "units", "fmt": "units", "f": "<O01>"},
 {"key": "wr", "id": "V04", "label": "Wholesale - retail units (proxy for channel stock change)", "unit": "units", "fmt": "units", "f": "[vs]-[del]", "src": "PROXY"},
 {"key": "asp", "id": "V05", "label": "ASP = automotive revenue per vehicle sold", "unit": "EUR k", "fmt": "k", "f": "[rev]*1000/[vs]", "bold": True, "fill": FILL_KEY,
  "src": "Porsche definition: automotive sales revenue / vehicles sold. Includes parts, services, non-vehicle revenue."},
 {"key": "aspp", "id": "V06", "label": "ASP published by Porsche", "unit": "EUR k", "fmt": "k", "f": 'IF(<O06>="","",<O06>)', "src": "O06: S005 p.20, S009 p.32, S011 p.10, S016 p.10"},
 {"key": "aspd", "id": "V07", "label": "Check: calculated - published (published rounded to EUR 1k)", "unit": "EUR k", "fmt": "k", "f": 'IF([aspp]="",0,[asp]-[aspp])'},
 {"key": "rpd", "id": "V08", "label": "Revenue per delivery (reference only; different denominator)", "unit": "EUR k", "fmt": "k", "f": "[rev]*1000/[del]"},
 {"key": "bev", "id": "V09", "label": "BEV share of deliveries (published)", "unit": "%", "fmt": "pct", "f": 'IF(<O07>="","",<O07>/100)'},
 {"kind": "sec", "label": "B. Year-on-year: volume effect vs ASP effect on automotive revenue"},
 {"key": "d_vs", "id": "V10", "label": "Change in vehicle sales", "unit": "%", "fmt": "pct", "f": "[vs]/[vs@p]-1"},
 {"key": "d_asp", "id": "V11", "label": "Change in ASP", "unit": "%", "fmt": "pct", "f": "[asp]/[asp@p]-1"},
 {"key": "d_rev", "id": "V12", "label": "Change in automotive revenue", "unit": "EUR m", "f": "[rev]-[rev@p]", "bold": True},
 {"key": "vol", "id": "V13", "label": "  Volume effect (unit change at prior-year ASP)", "unit": "EUR m", "f": "([vs]-[vs@p])*[asp@p]/1000"},
 {"key": "aspe", "id": "V14", "label": "  ASP effect (ASP change at current units)", "unit": "EUR m", "f": "([asp]-[asp@p])*[vs]/1000",
  "src": "Residual of price, model mix, derivative mix, options, FX and non-vehicle revenue - NOT 'price'"},
 {"key": "vchk", "id": "V15", "label": "Check: volume + ASP effect - revenue change (must be 0)", "unit": "EUR m", "fmt": "eur1", "f": "[vol]+[aspe]-[d_rev]"},
 {"kind": "sec", "label": "C. Delivery mix by model line (retail units; drives the ASP effect only indirectly)"},
]
for mid, mname in MODELS:
    CM.append({"key": f"m_{mid}", "id": "", "label": mname, "unit": "units", "fmt": "units",
               "f": f'SUMIFS({MXR("F")},{MXR("A")},{{c}}$5,{MXR("B")},"model",{MXR("C")},"{mid}")', "src": "Data_Mix (S025)"})
CM.append({"key": "m_tot", "id": "V16", "label": "Total by model", "unit": "units", "fmt": "units", "bold": True,
           "f": "+".join(f"[m_{m}]" for m, _ in MODELS)})
CM.append({"key": "m_chk", "id": "V17", "label": "Check: model total - deliveries (must be 0)", "unit": "units", "fmt": "units", "f": "[m_tot]-[del]"})
for mid, mname in MODELS:
    CM.append({"key": f"s_{mid}", "id": "", "label": f"  share {mname}", "unit": "%", "fmt": "pct", "f": f"[m_{mid}]/[m_tot]"})
for mid, mname in (("MAC_ICE", "Macan ICE"), ("MAC_BEV", "Macan BEV")):
    CM.append({"key": f"m_{mid}", "id": "", "label": f"  memo: {mname} (split published from FY2024)", "unit": "units", "fmt": "units",
               "f": f'IF(COUNTIFS({MXR("A")},{{c}}$5,{MXR("C")},"{mid}")=0,"",SUMIFS({MXR("F")},{MXR("A")},{{c}}$5,{MXR("C")},"{mid}"))'})
CM.append({"kind": "sec", "label": "D. Delivery mix by region"})
for rid, rname in REGIONS:
    CM.append({"key": f"r_{rid}", "id": "", "label": rname, "unit": "units", "fmt": "units",
               "f": f'SUMIFS({MXR("F")},{MXR("A")},{{c}}$5,{MXR("B")},"region",{MXR("C")},"{rid}")'})
CM.append({"key": "r_tot", "id": "V18", "label": "Total by region", "unit": "units", "fmt": "units", "bold": True,
           "f": "+".join(f"[r_{r}]" for r, _ in REGIONS)})
CM.append({"key": "r_chk", "id": "V19", "label": "Check: region total - deliveries (must be 0)", "unit": "units", "fmt": "units", "f": "[r_tot]-[del]"})
for rid, rname in REGIONS:
    CM.append({"key": f"rs_{rid}", "id": "", "label": f"  share {rname}", "unit": "%", "fmt": "pct", "f": f"[r_{rid}]/[r_tot]"})
CM_END = layout("Commercial", CM)

# ================================================================ Cash spec
ann = rd("annotation.csv")[1:]
cash_notes = {r[0]: f"{r[2]} [{r[3]} {r[4]}]" for r in ann if r[1] == "Cash flow"}
PUB_NCF = {"2024-H1": 0.063, "2025-H1": 0.024, "2026-H1": 0.067}
CS = [
 {"kind": "sec", "label": "A. Automotive profit vs cash"},
 {"key": "rev", "id": "K01", "label": "Automotive sales revenue", "unit": "EUR m", "f": "<A01>"},
 {"key": "ebit", "id": "K02", "label": "Automotive EBIT", "unit": "EUR m", "f": "<A02>"},
 {"key": "ebitda", "id": "K03", "label": "Automotive EBITDA", "unit": "EUR m", "f": "<A03>"},
 {"key": "cfo", "id": "K04", "label": "Cash flow from operating activities", "unit": "EUR m", "f": "<A13a>", "src": "S025 sheet 02"},
 {"key": "cfi", "id": "K05", "label": "Investing activities of current operations", "unit": "EUR m", "f": "<A13b>", "src": "S025 sheet 02"},
 {"key": "ncf", "id": "K06", "label": "Automotive net cash flow", "unit": "EUR m", "f": "<A13>", "bold": True, "fill": FILL_KEY, "src": "S025 sheet 02"},
 {"key": "ncfchk", "id": "K07", "label": "Check: NCF - (CFO + investing)", "unit": "EUR m", "fmt": "eur1", "f": "[ncf]-([cfo]+[cfi])", "src": "AF-12: H1 2024 NCF stored rounded (EUR 0.11m)"},
 {"key": "liq", "id": "K08", "label": "Automotive net liquidity (period end)", "unit": "EUR m", "f": "<A14>"},
 {"kind": "sec", "label": "B. Ratios and reconciliation to Porsche"},
 {"key": "ncfm", "id": "K09", "label": "Net cash flow margin (NCF / automotive revenue)", "unit": "%", "fmt": "pct", "f": "[ncf]/[rev]", "bold": True},
 {"key": "ncfm_pub", "id": "K10", "label": "Published NCF margin (reference)", "unit": "%", "fmt": "pct", "values": PUB_NCF,
  "src": "Input: S009 slide 35 (H1 2024), S016 slide 13 (H1 2025, H1 2026)"},
 {"key": "ncfm_d", "id": "K11", "label": "Check: calculated - published", "unit": "pp", "fmt": "pp", "f": 'IF([ncfm_pub]="",0,([ncfm]-[ncfm_pub])*100)'},
 {"key": "ebm", "id": "K12", "label": "EBITDA margin (calculated)", "unit": "%", "fmt": "pct", "f": "[ebitda]/[rev]"},
 {"key": "ebm_pub", "id": "K13", "label": "EBITDA margin (published)", "unit": "%", "fmt": "pct", "f": 'IF(<A04>="","",<A04>/100)'},
 {"key": "ebm_d", "id": "K14", "label": "Check: calculated - published", "unit": "pp", "fmt": "pp", "f": 'IF([ebm_pub]="",0,([ebm]-[ebm_pub])*100)'},
 {"key": "cfo_e", "id": "K15", "label": "Operating cash flow / EBITDA", "unit": "x", "fmt": "x", "f": "[cfo]/[ebitda]"},
 {"key": "ncf_e", "id": "K16", "label": "Net cash flow / automotive EBIT", "unit": "x", "fmt": "x", "f": 'IF([ebit]>0,[ncf]/[ebit],"n/m")'},
 {"kind": "sec", "label": "C. Simple cash proxy (EBITDA - capex - capitalised development costs) - PROXY"},
 {"key": "capex", "id": "K17", "label": "Automotive capex", "unit": "EUR m", "f": "<A10>"},
 {"key": "capdev", "id": "K18", "label": "Capitalised development costs (accounting amount used as investment proxy)", "unit": "EUR m", "f": "<A06>"},
 {"key": "inv", "id": "K19", "label": "Investment proxy (capex + capitalised development costs)", "unit": "EUR m", "f": "[capex]+[capdev]"},
 {"key": "othinv", "id": "K20", "label": "Other investing outflow (-investing - investment proxy)", "unit": "EUR m", "f": "-[cfi]-[inv]"},
 {"key": "proxy", "id": "K21", "label": "Simple cash proxy", "unit": "EUR m", "f": "[ebitda]-[capex]-[capdev]", "bold": True},
 {"key": "rem", "id": "K22", "label": "NCF - proxy (working capital, tax, provisions, other)", "unit": "EUR m", "f": "[ncf]-[proxy]", "bold": True},
]
CS_END = layout("Cash_Crosscheck", CS)

# ================================================================ render analysis sheets
sheet_frame("EBIT_Analysis", "Profitability / EBIT Analysis - reported vs adjusted, and the year-on-year bridge",
            "Group level. Headline levels L0/L1/L2. L3 is a supplementary hybrid sensitivity only (section E).",
            "After removing Porsche's disclosed exceptional items, is group profitability improving - and what moved EBIT?")
write_sheet("EBIT_Analysis", EB)
sheet_frame("RD_DA", "R&D & D&A Analysis - capitalisation, P&L charge and impairment treatment",
            "Automotive segment. Total R&D costs are Porsche's accounting definition (not cash).",
            "Are non-cash accounting effects (R&D capitalisation, depreciation, impairments) flattering or hurting EBIT?")
write_sheet("RD_DA", RDS)
sheet_frame("Commercial", "Commercial / Value-over-Volume Analysis",
            "Automotive segment. ASP uses Porsche's definition (automotive revenue / vehicles sold, wholesale).",
            "Is higher revenue per car compensating for lower volume?")
write_sheet("Commercial", CM)
sheet_frame("Cash_Crosscheck", "Cash Cross-check - does cash confirm the profit picture?",
            "Automotive segment. Net cash flow per Porsche definition (fact sheet S025).",
            "Does automotive cash generation move with EBIT/EBITDA, or do timing items dominate?")
write_sheet("Cash_Crosscheck", CS)
wsC = WS["Cash_Crosscheck"]
r = CS_END + 1
wsC.cell(row=r, column=2, value="Cash notes (verified, mart.annotation)").font = F_SEC
for j, p in enumerate(P):
    if p in cash_notes:
        c = wsC[f"{PC[j]}{r + 1}"]; c.value = cash_notes[p]; c.font = F_SUB; c.alignment = WR
wsC.row_dimensions[r + 1].height = 120

# ================================================================ SQL_Outputs
ws = WS["SQL_Outputs"]
title(ws, "SQL_Outputs - mart view results imported for reconciliation only",
      "Blue = value exported from the PostgreSQL mart views after Phase 4 cleanup. The analysis sheets do NOT use these values; "
      "Audit_Checks compares them with the Excel formulas.",
      {"A": 7, "B": 52, "C": 10, **{c: 12 for c in PC}, "K": 50})
so = rd("sqlout.csv"); cols = so[0]; data = {r[0]: r for r in so[1:]}
hdr(ws, HDR_ROW, ["", "mart column", "unit"] + P + ["mart view"])
SQLROW = {}
VIEWOF = {"ebit_l0_reported": "v_ebit_adjusted", "ebit_l1_ex_realignment": "v_ebit_adjusted", "ebit_l2_ex_realignment_tariffs": "v_ebit_adjusted",
          "ebit_l3_hybrid_sensitivity": "v_ebit_adjusted", "margin_l0_reported_pct": "v_ebit_adjusted", "margin_l1_pct": "v_ebit_adjusted",
          "margin_l2_pct": "v_ebit_adjusted", "margin_l3_hybrid_pct": "v_ebit_adjusted", "da_clean": "v_da_clean",
          "rd_pl_charge": "v_rd_capitalisation", "ebit_effect_of_cap_rate_change": "v_rd_capitalisation", "asp_calc_eur_k": "v_asp_volume",
          "volume_effect": "v_asp_volume", "asp_effect": "v_asp_volume", "ncf_margin_pct": "v_cash_crosscheck",
          "simple_cash_proxy": "v_cash_crosscheck", "eff_residual_other_drivers": "v_ebit_bridge_yoy", "pp_residual_other_drivers": "v_ebit_bridge_yoy"}
for i, cname in enumerate(cols[1:]):
    rr = FIRST + i; SQLROW[cname] = rr
    ws.cell(row=rr, column=2, value=cname).font = F_N
    ws.cell(row=rr, column=3, value="% x100" if "pct" in cname else ("EUR k" if "eur_k" in cname else ("pp" if cname.startswith("pp_") else "EUR m"))).font = F_N
    ws.cell(row=rr, column=11, value="mart." + VIEWOF[cname]).font = F_SUB
    for j, p in enumerate(P):
        v = data[p][i + 1]
        if v == "": continue
        c = ws[f"{PC[j]}{rr}"]; c.value = float(v); c.font = F_IN; c.number_format = '#,##0.000'
g = rd("gsum.csv"); gh, gv = g[0], g[1]
r0 = FIRST + len(cols) + 2
ws.cell(row=r0 - 1, column=2, value="mart.v_guidance_2026_summary (scenario combinations; % x100)").font = F_SEC
GSROW = {}
for i, (k, v) in enumerate(zip(gh, gv)):
    if k in ("scope", "scenario_type"): continue
    rr = r0 + len(GSROW); GSROW[k] = rr
    ws.cell(row=rr, column=2, value=k).font = F_N
    c = ws.cell(row=rr, column=4, value=float(v)); c.font = F_IN; c.number_format = '#,##0.000'
r1 = r0 + len(GSROW) + 2
ws.cell(row=r1 - 1, column=2, value="mart.v_guidance_2026 (27 scenario combinations; % x100)").font = F_SEC
hdr(ws, r1, ["", "scenario_id", "", "h2_margin_l0", "h2_margin_l1", "h2_margin_l2"])
GG = {}
for i, row in enumerate(rd("ggrid.csv")[1:]):
    rr = r1 + 1 + i; GG[row[0]] = rr
    ws.cell(row=rr, column=2, value=row[0]).font = F_N
    for j in range(3):
        c = ws.cell(row=rr, column=4 + j, value=float(row[1 + j])); c.font = F_IN; c.number_format = '0.00'
r2 = r1 + 1 + len(GG) + 2
ws.cell(row=r2 - 1, column=2, value="mart.v_guidance_2026_auto (9 scenario combinations; % x100)").font = F_SEC
hdr(ws, r2, ["", "scenario_id", "", "h2_auto_ebitda_margin", "h2_auto_ncf_margin"])
GA = {}
for i, row in enumerate(rd("gauto.csv")[1:]):
    rr = r2 + 1 + i; GA[row[0]] = rr
    ws.cell(row=rr, column=2, value=row[0]).font = F_N
    for j in range(2):
        c = ws.cell(row=rr, column=4 + j, value=float(row[1 + j])); c.font = F_IN; c.number_format = '0.00'
ws.freeze_panes = "D6"

# ================================================================ Guidance_Scenarios
ws = WS["Guidance_Scenarios"]
title(ws, "FY2026 Guidance Scenarios - what the guidance arithmetically implies for H2 2026",
      "All rows are SCENARIO COMBINATIONS of the ends of Porsche's guidance ranges. They are NOT forecasts and carry NO probabilities. "
      "'mid' = midpoint of every range, not a most-likely case.",
      {"A": 12, "B": 40, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12, "J": 12, "K": 12, "L": 12, "M": 12, "N": 12, "O": 12, "P": 40})
ws["A3"] = "Business question: If Porsche meets its FY2026 guidance, what must H2 2026 look like - reported and excluding exceptional items?"; ws["A3"].font = F_B
for col in range(1, 17):
    ws.cell(row=4, column=col).fill = FILL_SCEN
ws["A4"] = "SCENARIO COMBINATIONS - NOT FORECASTS - NO PROBABILITIES"; ws["A4"].font = F_WARN
# guidance inputs
hdr(ws, 6, ["guidance_id", "metric", "unit", "low", "mid", "high", "basis", "", "", "label", "source", "source ref"])
gin = rd("guidance.csv")[1:]
GROW = {}
for i, rr_ in enumerate(gin):
    rr = 7 + i; gid = rr_[0]; GROW[gid] = rr
    pct = rr_[5] == "%"
    ws.cell(row=rr, column=1, value=gid).font = F_N
    ws.cell(row=rr, column=2, value=rr_[1]).font = F_N
    ws.cell(row=rr, column=3, value=rr_[5]).font = F_N
    lo, hi = float(rr_[3]), float(rr_[4])
    for col, v in (("D", lo / 100 if pct else lo), ("F", hi / 100 if pct else hi)):
        c = ws[f"{col}{rr}"]; c.value = v; c.font = F_IN; c.number_format = FMT["pct"] if pct else FMT["eur"]
    ws[f"E{rr}"] = f"=(D{rr}+F{rr})/2"; ws[f"E{rr}"].number_format = FMT["pct2"] if pct else FMT["eur"]; ws[f"E{rr}"].font = F_N
    ws.cell(row=rr, column=7, value=rr_[6]).font = F_SUB
    ws.cell(row=rr, column=10, value=rr_[7]).font = F_N
    ws.cell(row=rr, column=11, value=rr_[8]).font = F_N
    ws.cell(row=rr, column=12, value=rr_[9]).font = F_SUB
ws["A14"] = "% ranges entered as fractions (5.5% = 0.055). GD_EXTRA comes from the CFO statement on the H1 2026 call (S027, Tier 3 record); consistent with S018 ('three-digit-million amount in H2')."
ws["A14"].font = F_SUB
# H1 actuals and assumptions
hdr(ws, 16, ["ref", "H1 2026 actuals and assumptions", "value", "", "", "", "lineage"])
ACT = [
 ("h1rev", "Group revenue H1 2026", "=EBIT_Analysis!J{rev}", "eur", "EBIT_Analysis E01"),
 ("h1ebit", "Reported EBIT H1 2026 (L0)", "=EBIT_Analysis!J{l0}", "eur", "EBIT_Analysis E02"),
 ("h1exc", "Exceptional items H1 2026 (net, X01)", "=EBIT_Analysis!J{x01}", "eur", "EBIT_Analysis E06"),
 ("h1tar", "US tariffs H1 2026 (X03)", "=EBIT_Analysis!J{x03}", "eur", "EBIT_Analysis E07"),
 ("h1m0", "Reported margin H1 2026", "=EBIT_Analysis!J{m0}", "pct", "EBIT_Analysis E03"),
 ("h1m1", "L1 margin H1 2026", "=EBIT_Analysis!J{m1}", "pct", "EBIT_Analysis E09"),
 ("h1m2", "L2 margin H1 2026", "=EBIT_Analysis!J{m2}", "pct", "EBIT_Analysis E11"),
 ("h2m0", "Reported margin H2 2025 (actual, DERIVED)", "=EBIT_Analysis!I{m0}", "pct", "EBIT_Analysis E03"),
 ("h2m1", "L1 margin H2 2025 (actual, DERIVED)", "=EBIT_Analysis!I{m1}", "pct", "EBIT_Analysis E09"),
 ("h1arev", "Automotive revenue H1 2026", "=Cash_Crosscheck!J{crev}", "eur", "Cash_Crosscheck K01"),
 ("h1aeb", "Automotive EBITDA H1 2026", "=Cash_Crosscheck!J{cebitda}", "eur", "Cash_Crosscheck K03"),
 ("h1ancf", "Automotive net cash flow H1 2026", "=Cash_Crosscheck!J{cncf}", "eur", "Cash_Crosscheck K06"),
 ("share", "A-05: automotive share of group revenue (H1 2026)", "=C{h1arev}/C{h1rev}", "pct2", "Assumption A-05"),
 ("h2tar", "A-07: H2 2026 tariffs = H1 2026 tariffs", "=C{h1tar}", "eur", "Assumption A-07 (base)"),
 ("h2tar_alt", "A-07 sensitivity: H2 2026 tariffs (secondary FY ~700m -> H2 ~300m)", 300, "eur", "Input: Quartr summary of Q1 2026 call (secondary)"),
]
AROW = {}
for i, (k, lab, f, fm, lin) in enumerate(ACT):
    AROW[k] = 17 + i
for i, (k, lab, f, fm, lin) in enumerate(ACT):
    rr = AROW[k]
    ws.cell(row=rr, column=1, value=k).font = F_N
    ws.cell(row=rr, column=2, value=lab).font = F_N
    if isinstance(f, str):
        f = f.format(rev=REG[("EBIT_Analysis", "rev")], l0=REG[("EBIT_Analysis", "l0")], x01=REG[("EBIT_Analysis", "x01")],
                     x03=REG[("EBIT_Analysis", "x03")], m0=REG[("EBIT_Analysis", "m0")], m1=REG[("EBIT_Analysis", "m1")],
                     m2=REG[("EBIT_Analysis", "m2")], crev=REG[("Cash_Crosscheck", "rev")], cebitda=REG[("Cash_Crosscheck", "ebitda")],
                     cncf=REG[("Cash_Crosscheck", "ncf")], h1arev=AROW["h1arev"], h1rev=AROW["h1rev"], h1tar=AROW["h1tar"])
        c = ws.cell(row=rr, column=3, value=f); c.font = F_LINK if "!" in f else F_N
    else:
        c = ws.cell(row=rr, column=3, value=f); c.font = F_IN
    c.number_format = FMT[fm]
    ws.cell(row=rr, column=7, value=lin).font = F_SUB
A = lambda k: f"$C${AROW[k]}"
# selected scenario
SEL = 34
ws.cell(row=SEL - 1, column=1, value="Selected scenario combination (choose low / mid / high)").font = F_SEC
for col in range(1, 17): ws.cell(row=SEL - 1, column=col).fill = FILL_SEC
dv = DataValidation(type="list", formula1='"low,mid,high"', allow_blank=False); ws.add_data_validation(dv)
sel_inputs = [("Revenue range end", "GD_REV"), ("Return-on-sales range end", "GD_ROS"), ("Extraordinary-expense range end", "GD_EXTRA")]
for i, (lab, gid) in enumerate(sel_inputs):
    rr = SEL + i
    ws.cell(row=rr, column=2, value=lab).font = F_N
    c = ws.cell(row=rr, column=3, value="mid"); c.font = F_IN; c.fill = FILL_KEY; c.alignment = CEN; dv.add(c)
    ws.cell(row=rr, column=4, value=f"=SUMIFS($D${GROW[gid]}:$F${GROW[gid]},$D$6:$F$6,C{rr})").number_format = FMT["pct2"] if gid == "GD_ROS" else FMT["eur"]
SELOUT = [
 ("sel_fyebit", "FY2026 EBIT implied (revenue x RoS)", f"=D{SEL}*D{SEL+1}", "eur"),
 ("sel_h2rev", "H2 2026 revenue implied", f"=D{SEL}-{A('h1rev')}", "eur"),
 ("sel_h2ebit", "H2 2026 reported EBIT implied", f"=D{SEL+3}-{A('h1ebit')}", "eur"),
 ("sel_h2m0", "H2 2026 reported margin implied", f"=D{SEL+5}/D{SEL+4}", "pct"),
 ("sel_h2exc", "H2 2026 exceptional items implied (FY - H1 net)", f"=D{SEL+2}-{A('h1exc')}", "eur"),
 ("sel_h2m1", "H2 2026 L1 margin implied", f"=(D{SEL+5}+D{SEL+7})/D{SEL+4}", "pct"),
 ("sel_h2m2", "H2 2026 L2 margin implied (A-07 base)", f"=(D{SEL+5}+D{SEL+7}+{A('h2tar')})/D{SEL+4}", "pct"),
 ("sel_h2m2a", "H2 2026 L2 margin implied (A-07 sensitivity)", f"=(D{SEL+5}+D{SEL+7}+{A('h2tar_alt')})/D{SEL+4}", "pct"),
 ("sel_cmp", "Compare: H1 2026 L1 margin (actual)", f"={A('h1m1')}", "pct"),
]
for i, (k, lab, f, fm) in enumerate(SELOUT):
    rr = SEL + 3 + i
    ws.cell(row=rr, column=2, value=lab).font = F_B if k in ("sel_h2m0", "sel_h2m1") else F_N
    c = ws.cell(row=rr, column=4, value=f); c.number_format = FMT[fm]; c.font = F_N
    if k in ("sel_h2m0", "sel_h2m1"): c.fill = FILL_KEY
# full grid
GRID = SEL + 14
ws.cell(row=GRID - 1, column=1, value="All 27 scenario combinations (group) - arithmetic only").font = F_SEC
for col in range(1, 17): ws.cell(row=GRID - 1, column=col).fill = FILL_SEC
gh = ["scenario_id", "scenario type", "revenue", "RoS", "extraord.", "FY revenue", "FY RoS", "FY extraord.", "FY EBIT", "H2 revenue",
      "H2 EBIT L0", "H2 margin L0", "H2 exc. implied", "H2 EBIT L1", "H2 margin L1", "H2 margin L2"]
hdr(ws, GRID, gh)
cases = ["low", "mid", "high"]
gi = 0
for cr in cases:
    for cs in cases:
        for ce in cases:
            rr = GRID + 1 + gi; gi += 1
            vals = [f"{cr}/{cs}/{ce}", SCEN_LABEL, cr, cs, ce]
            for j, v in enumerate(vals):
                ws.cell(row=rr, column=1 + j, value=v).font = F_N
            fx = {
              "F": f"=SUMIFS($D${GROW['GD_REV']}:$F${GROW['GD_REV']},$D$6:$F$6,C{rr})",
              "G": f"=SUMIFS($D${GROW['GD_ROS']}:$F${GROW['GD_ROS']},$D$6:$F$6,D{rr})",
              "H": f"=SUMIFS($D${GROW['GD_EXTRA']}:$F${GROW['GD_EXTRA']},$D$6:$F$6,E{rr})",
              "I": f"=F{rr}*G{rr}",
              "J": f"=F{rr}-{A('h1rev')}",
              "K": f"=I{rr}-{A('h1ebit')}",
              "L": f"=K{rr}/J{rr}",
              "M": f"=H{rr}-{A('h1exc')}",
              "N": f"=K{rr}+M{rr}",
              "O": f"=N{rr}/J{rr}",
              "P": f"=(N{rr}+{A('h2tar')})/J{rr}",
            }
            for col, f in fx.items():
                c = ws[f"{col}{rr}"]; c.value = f; c.font = F_N
                c.number_format = FMT["pct"] if col in ("G", "L", "O", "P") else FMT["eur"]
            ws[f"B{rr}"].font = F_SUB
GRID_L = GRID + 27
SUMR = GRID_L + 3
ws.cell(row=SUMR - 1, column=1, value="Summary across the 27 scenario combinations (min / all-midpoints / max) vs actuals").font = F_SEC
for col in range(1, 17): ws.cell(row=SUMR - 1, column=col).fill = FILL_SEC
hdr(ws, SUMR, ["", "H2 2026 implied", "min", "all midpoints", "max", "H1 2026 actual", "H2 2025 actual"])
midrow = GRID + 1 + 13     # mid/mid/mid is the 14th combination
SUMS = [("sum_m0", "Reported margin (L0)", "L", A("h1m0"), A("h2m0")),
        ("sum_m1", "L1 margin (excl. realignment & battery)", "O", A("h1m1"), A("h2m1")),
        ("sum_m2", "L2 margin (also excl. tariffs, A-07)", "P", A("h1m2"), '=EBIT_Analysis!I' + str(REG[("EBIT_Analysis", "m2")])),
        ("sum_rev", "H2 revenue (EUR m)", "J", None, None),
        ("sum_ebit", "H2 reported EBIT (EUR m)", "K", None, None)]
SUMROW = {}
for i, (k, lab, col, h1, h2) in enumerate(SUMS):
    rr = SUMR + 1 + i; SUMROW[k] = rr
    ws.cell(row=rr, column=2, value=lab).font = F_B if k in ("sum_m0", "sum_m1") else F_N
    fm = FMT["pct"] if col in ("L", "O", "P") else FMT["eur"]
    for cc, f in (("C", f"=MIN({col}{GRID+1}:{col}{GRID_L})"), ("D", f"={col}{midrow}"), ("E", f"=MAX({col}{GRID+1}:{col}{GRID_L})")):
        c = ws[f"{cc}{rr}"]; c.value = f; c.number_format = fm; c.font = F_N
    if h1:
        ws[f"F{rr}"] = "=" + h1.lstrip("="); ws[f"F{rr}"].number_format = fm; ws[f"F{rr}"].font = F_N
        ws[f"G{rr}"] = h2 if h2.startswith("=") else "=" + h2; ws[f"G{rr}"].number_format = fm
        ws[f"G{rr}"].font = F_LINK if "!" in ws[f"G{rr}"].value else F_N
ws[f"B{SUMR+6}"] = ("Reading: the midpoint column combines all guidance midpoints. A scenario range wider than the guidance itself is expected "
                    "because the corners are combined independently (assumption A-08).")
ws[f"B{SUMR+6}"].font = F_SUB
# automotive grid
AG = SUMR + 9
ws.cell(row=AG - 1, column=1, value="Automotive: 9 scenario combinations (revenue end x margin end) - arithmetic only").font = F_SEC
for col in range(1, 17): ws.cell(row=AG - 1, column=col).fill = FILL_SEC
hdr(ws, AG, ["scenario_id", "scenario type", "revenue", "margins", "", "FY group rev", "A-05 share", "FY auto rev", "H2 auto rev",
             "FY EBITDA m.", "H2 EBITDA m.", "FY NCF m.", "H2 NCF m.", "H1 26 EBITDA m.", "H1 26 NCF m."])
ai = 0
for cr in cases:
    for cm_ in cases:
        rr = AG + 1 + ai; ai += 1
        for j, v in enumerate([f"{cr}/{cm_}", SCEN_LABEL, cr, cm_]):
            ws.cell(row=rr, column=1 + j, value=v).font = F_N
        ws[f"B{rr}"].font = F_SUB
        fx = {
          "F": f"=SUMIFS($D${GROW['GD_REV']}:$F${GROW['GD_REV']},$D$6:$F$6,C{rr})",
          "G": f"={A('share')}",
          "H": f"=F{rr}*G{rr}",
          "I": f"=H{rr}-{A('h1arev')}",
          "J": f"=SUMIFS($D${GROW['GD_EBITDA']}:$F${GROW['GD_EBITDA']},$D$6:$F$6,D{rr})",
          "K": f"=(J{rr}*H{rr}-{A('h1aeb')})/I{rr}",
          "L": f"=SUMIFS($D${GROW['GD_NCF']}:$F${GROW['GD_NCF']},$D$6:$F$6,D{rr})",
          "M": f"=(L{rr}*H{rr}-{A('h1ancf')})/I{rr}",
          "N": f"={A('h1aeb')}/{A('h1arev')}",
          "O": f"={A('h1ancf')}/{A('h1arev')}",
        }
        for col, f in fx.items():
            c = ws[f"{col}{rr}"]; c.value = f; c.font = F_N
            c.number_format = FMT["pct2"] if col == "G" else (FMT["pct"] if col in ("J", "K", "L", "M", "N", "O") else FMT["eur"])
AG_L = AG + 9
ws.freeze_panes = "C6"

# ================================================================ KPI_Overview
ws = WS["KPI_Overview"]
title(ws, "KPI Overview - Porsche 2026 margin recovery",
      "Half-year view, group and automotive. All values are formulas linked to the analysis sheets (green = link).",
      {"A": 7, "B": 46, "C": 8, **{c: 12 for c in PC}, "K": 3, "L": 12, "M": 12, "N": 12, "O": 48})
ws["A3"] = ("Question: Porsche's reported margin rose from 5.5% to 7.8% in H1 2026. After removing the exceptional items Porsche discloses, "
            "and isolating non-cash effects, is underlying profitability improving - and what does H2 2026 guidance imply?")
ws["A3"].font = F_B; ws.merge_cells("A3:O3"); ws["A3"].alignment = WR; ws.row_dimensions[3].height = 30
hdr(ws, HDR_ROW, ["", "KPI", "unit"] + [None] * 7 + ["", "H1 2025", "H1 2026", "change", "definition / lineage"])
for c in PC:
    ws[f"{c}{HDR_ROW}"] = f"=Data_Panel!{c}${HDR_ROW}"; ws[f"{c}{HDR_ROW}"].font = F_H
KPIS = [
 ("Group sales revenue", "EUR m", "EBIT_Analysis", "rev", "eur", "abs", "Group, fact sheet S025"),
 ("L0 Reported EBIT", "EUR m", "EBIT_Analysis", "l0", "eur", "abs", "Group operating profit"),
 ("L0 Reported margin", "%", "EBIT_Analysis", "m0", "pct", "pp", "EBIT / revenue"),
 ("L1 EBIT excl. realignment & battery", "EUR m", "EBIT_Analysis", "l1", "eur", "abs", "Headline adjusted level"),
 ("L1 margin", "%", "EBIT_Analysis", "m1", "pct", "pp", "Headline adjusted level"),
 ("L2 EBIT excl. realignment, battery & tariffs", "EUR m", "EBIT_Analysis", "l2", "eur", "abs", "Headline adjusted level"),
 ("L2 margin", "%", "EBIT_Analysis", "m2", "pct", "pp", "Headline adjusted level"),
 ("Automotive revenue", "EUR m", "Commercial", "rev", "eur", "abs", "Automotive segment"),
 ("Vehicle sales (wholesale)", "units", "Commercial", "vs", "units", "rel", "Fact sheet S025"),
 ("ASP (automotive revenue per vehicle sold)", "EUR k", "Commercial", "asp", "k", "rel", "Porsche ASP definition"),
 ("Deliveries (retail)", "units", "Commercial", "del", "units", "rel", "Fact sheet S025"),
 ("Automotive net cash flow", "EUR m", "Cash_Crosscheck", "ncf", "eur", "abs", "Fact sheet S025"),
 ("Net cash flow margin", "%", "Cash_Crosscheck", "ncfm", "pct", "pp", "NCF / automotive revenue"),
]
KROW = {}
for i, (lab, unit, sh, key, fm, chg, lin) in enumerate(KPIS):
    rr = FIRST + i; KROW[key + sh] = rr
    ws.cell(row=rr, column=2, value=lab).font = F_B if key in ("m0", "m1", "m2") else F_N
    ws.cell(row=rr, column=3, value=unit).font = F_N
    for c in PC:
        cell = ws[f"{c}{rr}"]; cell.value = f"={sh}!{c}{REG[(sh, key)]}"; cell.font = F_LINK; cell.number_format = FMT[fm]; cell.border = BD
        if key in ("m1", "m2", "l1", "l2"): cell.fill = FILL_KEY
    ws[f"L{rr}"] = f"=H{rr}"; ws[f"M{rr}"] = f"=J{rr}"
    for cc in ("L", "M"): ws[f"{cc}{rr}"].number_format = FMT[fm]; ws[f"{cc}{rr}"].font = F_N; ws[f"{cc}{rr}"].border = BD
    ws[f"N{rr}"] = (f"=(M{rr}-L{rr})*100" if chg == "pp" else (f"=M{rr}/L{rr}-1" if chg == "rel" else f"=M{rr}-L{rr}"))
    ws[f"N{rr}"].number_format = FMT["pp"] if chg == "pp" else (FMT["pct"] if chg == "rel" else FMT["eur"]); ws[f"N{rr}"].font = F_B; ws[f"N{rr}"].border = BD
    ws[f"O{rr}"] = lin; ws[f"O{rr}"].font = F_SUB
KP_L = FIRST + len(KPIS) - 1
ws[f"B{KP_L+1}"] = "L3 (supplementary hybrid sensitivity) is deliberately not shown here; see EBIT_Analysis section E."; ws[f"B{KP_L+1}"].font = F_SUB
# readings
RR = KP_L + 3
ws.cell(row=RR, column=1, value="Formula-driven readings (descriptive; interpretation follows in Phase 6)").font = F_SEC
for col in range(1, 16): ws.cell(row=RR, column=col).fill = FILL_SEC
E = lambda k, c="J": f"EBIT_Analysis!{c}{REG[('EBIT_Analysis', k)]}"
S = lambda k, c="C": f"Guidance_Scenarios!{c}{SUMROW[k]}"
readings = [
 f'="1. Reported margin: "&TEXT({E("m0","H")},"0.0%")&" (H1 2025) -> "&TEXT({E("m0")},"0.0%")&" (H1 2026), "&TEXT({E("d_m0")},"+0.0;-0.0")&" pp; reported EBIT "&TEXT({E("d_l0")},"+#,##0;-#,##0")&" EUR m."',
 f'="2. Excluding disclosed realignment & battery items (L1): "&TEXT({E("m1","H")},"0.0%")&" -> "&TEXT({E("m1")},"0.0%")&" ("&TEXT({E("d_m1")},"+0.0;-0.0")&" pp); L1 EBIT "&TEXT({E("d_l1")},"+#,##0;-#,##0")&" EUR m."',
 f'="3. Also excluding US tariffs (L2): "&TEXT({E("m2","H")},"0.0%")&" -> "&TEXT({E("m2")},"0.0%")&" ("&TEXT({E("d_m2")},"+0.0;-0.0")&" pp)."',
 f'="4. Of the "&TEXT({E("b_d")},"+#,##0;-#,##0")&" EUR m change in reported EBIT, lower exceptional items contributed "&TEXT({E("e_exc")},"+#,##0;-#,##0")&" EUR m; residual / other EBIT drivers "&TEXT({E("e_res")},"+#,##0;-#,##0")&" EUR m."',
 f'="5. ASP "&TEXT(Commercial!J{REG[("Commercial","asp")]},"0.0")&" EUR k ("&TEXT(Commercial!J{REG[("Commercial","d_asp")]},"+0.0%;-0.0%")&"), vehicle sales "&TEXT(Commercial!J{REG[("Commercial","d_vs")]},"+0.0%;-0.0%")&": volume effect "&TEXT(Commercial!J{REG[("Commercial","vol")]},"+#,##0;-#,##0")&" vs ASP effect "&TEXT(Commercial!J{REG[("Commercial","aspe")]},"+#,##0;-#,##0")&" EUR m."',
 f'="6. Guidance scenario combinations (not forecasts) imply an H2 2026 reported margin of "&TEXT({S("sum_m0")},"0.0%")&"-"&TEXT({S("sum_m0","E")},"0.0%")&" (all midpoints "&TEXT({S("sum_m0","D")},"0.0%")&") and an L1 margin of "&TEXT({S("sum_m1")},"0.0%")&"-"&TEXT({S("sum_m1","E")},"0.0%")&" (midpoints "&TEXT({S("sum_m1","D")},"0.0%")&") vs "&TEXT({E("m1")},"0.0%")&" in H1 2026."',
 '="Model status: "&MODEL_STATUS_PLACEHOLDER',
]
for i, f in enumerate(readings):
    c = ws.cell(row=RR + 1 + i, column=2, value=f); c.font = F_N
    ws.merge_cells(start_row=RR + 1 + i, start_column=2, end_row=RR + 1 + i, end_column=15)
    c.alignment = Alignment(wrap_text=True, vertical="center"); ws.row_dimensions[RR + 1 + i].height = 28
# chart data for H1 2026 bridge
BR0 = RR + 10
ws.cell(row=BR0, column=2, value="H1 2026 vs H1 2025 EBIT bridge (EUR m) - chart data").font = F_SEC
brow = [("Exceptional items", "e_exc"), ("US tariffs", "e_tar"), ("Capitalised dev. costs", "e_rdc"), ("Clean D&A", "e_da"),
        ("Financial Services", "e_fs"), ("Residual / other drivers", "e_res"), ("Total change in EBIT", "b_d")]
for i, (lab, k) in enumerate(brow):
    ws.cell(row=BR0 + 1 + i, column=2, value=lab).font = F_N
    c = ws.cell(row=BR0 + 1 + i, column=4, value=f"={E(k)}"); c.font = F_LINK; c.number_format = FMT["eur"]
# charts
ch = LineChart(); ch.title = "Group margin by half-year: reported vs adjusted "
ch.y_axis.title = "margin"; ch.y_axis.number_format = '0%'; ch.height = 8; ch.width = 22; ch.legend.position = "b"
for key in ("m0", "m1", "m2"):
    rr = KROW[key + "EBIT_Analysis"]
    ref = Reference(ws, min_col=4, max_col=10, min_row=rr, max_row=rr)
    ch.add_data(ref, from_rows=True, titles_from_data=False)
for s, name in zip(ch.series, ["L0 reported", "L1 excl. realignment & battery", "L2 also excl. tariffs"]):
    from openpyxl.chart.series import SeriesLabel
    s.tx = SeriesLabel(v=name); s.smooth = False; s.marker.symbol = "circle"
ch.set_categories(Reference(ws, min_col=4, max_col=10, min_row=HDR_ROW, max_row=HDR_ROW))
ws.add_chart(ch, f"B{BR0 + 10}")
bc = BarChart(); bc.type = "bar"; bc.title = "H1 2026 vs H1 2025: what moved reported EBIT (EUR m)"; bc.x_axis.tickLblPos = "low"
bc.height = 8; bc.width = 18; bc.legend = None
bc.add_data(Reference(ws, min_col=4, min_row=BR0 + 1, max_row=BR0 + 7), titles_from_data=False)
bc.set_categories(Reference(ws, min_col=2, min_row=BR0 + 1, max_row=BR0 + 7))
ws.add_chart(bc, f"H{BR0 + 10}")
ws.freeze_panes = "D6"

# ================================================================ Audit_Register
ws = WS["Audit_Register"]
title(ws, "Audit_Register - sources, assumptions, known gaps, audit findings, SQL documentation",
      "Imported from the PostgreSQL audit and mart schemas (controlled text; do not edit).",
      {"A": 10, "B": 40, "C": 40, "D": 40, "E": 40, "F": 40, "G": 40, "H": 50})
r = 4
def block(ws, r, titletxt, header, data_rows):
    ws.cell(row=r, column=1, value=titletxt).font = F_SEC
    for col in range(1, 9): ws.cell(row=r, column=col).fill = FILL_SEC
    hdr(ws, r + 1, header)
    for i, row in enumerate(data_rows):
        for j, v in enumerate(row):
            c = ws.cell(row=r + 2 + i, column=j + 1, value=v); c.font = F_N; c.alignment = WR
    return r + 2 + len(data_rows) + 2
src = rd("sources.csv"); r = block(ws, r, "1. Source register (audit.dim_source)", src[0], src[1:])
asm = rd("assumption.csv"); ASM_TOP = r + 2; r = block(ws, r, "2. Assumption register (mart.assumption)", asm[0], asm[1:])
gap = rd("gaps.csv"); r = block(ws, r, "3. Known data gaps - looked for, not verifiable, NOT estimated (audit.data_gap)", gap[0], gap[1:])
fnd = rd("findings.csv"); r = block(ws, r, "4. Audit findings (audit.audit_finding)", fnd[0], fnd[1:])
cat = rd("catalog.csv"); r = block(ws, r, "5. SQL view catalogue (mart.view_catalog)", cat[0], cat[1:])
dq = rd("dq.csv"); DQ_TOP = r + 2; r = block(ws, r, "6. SQL data-quality results (audit.dq_result, summarised)", dq[0], [[x[0], x[1], int(x[2]), int(x[3]), int(x[4])] for x in dq[1:]])
DQ_BOT = DQ_TOP + len(dq) - 2
N_SRC = len(src) - 1

# ================================================================ Audit_Checks
ws = WS["Audit_Checks"]
title(ws, "Audit_Checks - model integrity and reconciliation to SQL",
      "Every check is a formula. PASS = within tolerance. Tolerances reflect source rounding.",
      {"A": 7, "B": 70, "C": 18, "D": 14, "E": 14, "F": 14, "G": 12, "H": 10, "I": 50})
ws["B4"] = "Overall model status"; ws["B4"].font = F_B
hdr(ws, 6, ["ID", "Check", "Category", "Expected", "Actual", "Difference", "Tolerance", "Status", "What it proves"])
ref = lambda sh, k, c1="D", c2="J": f"{sh}!{c1}{REG[(sh, k)]}:{c2}{REG[(sh, k)]}"
sq = lambda col, c1="D", c2="J": f"SQL_Outputs!{c1}{SQLROW[col]}:{c2}{SQLROW[col]}"
CH = []

# ---- portable scalar helpers: one term per cell, N() turns "" into 0, no array arithmetic ----
from openpyxl.utils import column_index_from_string as _ci, get_column_letter as _cl
def _cells(r):
    sh, rng = r.split("!"); a, b = rng.split(":") if ":" in rng else (rng, rng)
    ca, ra = re.match(r"\$?([A-Z]+)\$?(\d+)", a).groups(); cb, rb = re.match(r"\$?([A-Z]+)\$?(\d+)", b).groups()
    return [f"{sh}!{_cl(c)}{row}" for row in range(int(ra), int(rb) + 1) for c in range(_ci(ca), _ci(cb) + 1)]
def S_ABS(r):                     # sum of |x|
    return "+".join(f"ABS(N({x}))" for x in _cells(r))
def S_DIFF(ra, rb, mult=1):       # sum of |a*mult - b|
    return "+".join(f"ABS(N({x})*{mult}-N({y}))" for x, y in zip(_cells(ra), _cells(rb)))
def S_MAXABS(r):                  # max |x|
    return "MAX(" + ",".join(f"ABS(N({x}))" for x in _cells(r)) + ")"
def S_COUNT_GT(ra, rb):           # number of cells with |a| > b
    return "+".join(f"IF(ABS(N({x}))>N({y}),1,0)" for x, y in zip(_cells(ra), _cells(rb)))
def S_NOT_EQ(r, text):            # number of cells whose text differs from the exact label
    return "+".join(f'IF({x}="{text}",0,1)' for x in _cells(r))
def chk(cid, text, cat_, expected, actual, tol, proves):
    CH.append((cid, text, cat_, expected, actual, tol, proves))
chk("C01", "Data_Long rows imported = rows exported from SQL (236)", "Import", "236", f"COUNTA(Data_Long!B{DL_FIRST}:B{DL_LAST})", 0, "Complete import")
chk("C02", "Data_Panel holds every imported value of the panel variables", "Import",
    f'COUNTA(Data_Long!B{DL_FIRST}:B{DL_LAST})-COUNTIF(Data_Long!B{DL_FIRST}:B{DL_LAST},"G09")-COUNTIF(Data_Long!B{DL_FIRST}:B{DL_LAST},"G10")',
    f'COUNT(Data_Panel!D{FIRST}:J{PANEL_LAST})', 0, "Pivot loses nothing")
chk("C03", "Source register rows imported", "Import", str(N_SRC), f"COUNTA(Audit_Register!A6:A{5+N_SRC})", 0, "Lineage table complete")
recon = [
 ("C04", "L0 reported EBIT = SQL v_ebit_adjusted (sum of abs. differences, 7 periods)", ref("EBIT_Analysis", "l0"), sq("ebit_l0_reported"), 1, 0.01),
 ("C05", "L1 EBIT = SQL", ref("EBIT_Analysis", "l1"), sq("ebit_l1_ex_realignment"), 1, 0.01),
 ("C06", "L2 EBIT = SQL", ref("EBIT_Analysis", "l2"), sq("ebit_l2_ex_realignment_tariffs"), 1, 0.01),
 ("C07", "L3 hybrid sensitivity = SQL", ref("EBIT_Analysis", "l3"), sq("ebit_l3_hybrid_sensitivity"), 1, 0.01),
 ("C08", "L0 margin = SQL (SQL rounded to 0.01pp)", ref("EBIT_Analysis", "m0"), sq("margin_l0_reported_pct"), 100, 0.04),
 ("C09", "L1 margin = SQL", ref("EBIT_Analysis", "m1"), sq("margin_l1_pct"), 100, 0.04),
 ("C10", "L2 margin = SQL", ref("EBIT_Analysis", "m2"), sq("margin_l2_pct"), 100, 0.04),
 ("C11", "Clean D&A = SQL v_da_clean", ref("RD_DA", "daclean"), sq("da_clean"), 1, 0.01),
 ("C12", "R&D P&L charge = SQL v_rd_capitalisation", ref("RD_DA", "pl"), sq("rd_pl_charge"), 1, 0.01),
 ("C13", "Capitalisation-rate effect = SQL (5 comparable periods; SQL rounded to 0.1)", ref("RD_DA", "cfe", "F"), sq("ebit_effect_of_cap_rate_change", "F"), 1, 0.3),
 ("C14", "ASP = SQL v_asp_volume (SQL rounded to 0.01)", ref("Commercial", "asp"), sq("asp_calc_eur_k"), 1, 0.04),
 ("C15", "Volume effect = SQL (SQL rounded to 0.1)", ref("Commercial", "vol", "F"), sq("volume_effect", "F"), 1, 0.3),
 ("C16", "ASP effect = SQL", ref("Commercial", "aspe", "F"), sq("asp_effect", "F"), 1, 0.3),
 ("C17", "NCF margin = SQL v_cash_crosscheck", ref("Cash_Crosscheck", "ncfm"), sq("ncf_margin_pct"), 100, 0.04),
 ("C18", "Simple cash proxy = SQL", ref("Cash_Crosscheck", "proxy"), sq("simple_cash_proxy"), 1, 0.01),
 ("C19", "Bridge residual / other drivers = SQL v_ebit_bridge_yoy", ref("EBIT_Analysis", "e_res", "F"), sq("eff_residual_other_drivers", "F"), 1, 0.01),
 ("C20", "Bridge residual in pp = SQL (SQL rounded to 0.001pp)", ref("EBIT_Analysis", "p_res", "F"), sq("pp_residual_other_drivers", "F"), 1, 0.005),
]
for cid, text, a, b, mult, tol in recon:
    chk(cid, text, "Excel vs SQL", "0", S_DIFF(a, b, mult), tol, "Excel formulas reproduce the validated SQL mart")
g = "Guidance_Scenarios"
chk("C21", "Guidance summary (L0/L1/L2 min, mid, max) = SQL v_guidance_2026_summary", "Excel vs SQL", "0",
    "+".join(f"ABS(N({g}!{c}{SUMROW[k]})*100-N(SQL_Outputs!D{GSROW[s]}))" for k, pre in (("sum_m0", "h2_margin_l0"), ("sum_m1", "h2_margin_l1"), ("sum_m2", "h2_margin_l2"))
             for c, s in (("C", pre + "_min"), ("D", pre + "_mid"), ("E", pre + "_max"))), 0.05, "Scenario arithmetic reproduces SQL")
grid_terms = []
gi = 0
for cr in cases:
    for cs in cases:
        for ce in cases:
            rr = GRID + 1 + gi; gi += 1; sid = f"{cr}/{cs}/{ce}"
            for col, sc in (("L", "D"), ("O", "E"), ("P", "F")):
                grid_terms.append(f"ABS(N({g}!{col}{rr})*100-N(SQL_Outputs!{sc}{GG[sid]}))")
chk("C22", "All 27 guidance scenario combinations (L0/L1/L2 H2 margins) = SQL v_guidance_2026", "Excel vs SQL", "0", "+".join(grid_terms), 0.41, "Every grid row reproduces SQL")
auto_terms = []
ai = 0
for cr in cases:
    for cm_ in cases:
        rr = AG + 1 + ai; ai += 1; sid = f"{cr}/{cm_}"
        for col, sc in (("K", "D"), ("M", "E")):
            auto_terms.append(f"ABS(N({g}!{col}{rr})*100-N(SQL_Outputs!{sc}{GA[sid]}))")
chk("C23", "Automotive guidance combinations = SQL v_guidance_2026_auto", "Excel vs SQL", "0", "+".join(auto_terms), 0.1, "Automotive scenario arithmetic reproduces SQL")
chk("C24", "R&D identity: total - capitalised - expensed (sum of abs., EUR 1m rounding)", "Identity", "0", S_ABS(ref('RD_DA','chk_rd')), 7, "R&D table internally consistent")
chk("C25", "D&A identity: total - amortisation - depreciation", "Identity", "0", S_ABS(ref('RD_DA','chk_da')), 7, "D&A table internally consistent")
chk("C26", "EBITDA = automotive EBIT + D&A", "Identity", "0",
    "+".join(f"ABS(N({x})-N({y})-N({z}))" for x, y, z in zip(_cells(ref('Cash_Crosscheck','ebitda')), _cells(ref('Cash_Crosscheck','ebit')), _cells(ref('RD_DA','da')))), 7, "Segment EBITDA consistent with D&A")
chk("C27", "Margin-point bridge closes (components - total)", "Identity", "0", S_ABS(ref('EBIT_Analysis','p_chk','F')), 0.0001, "pp bridge exact")
chk("C28", "Reported margin = Porsche published return on sales (0.1pp rounding)", "Porsche reconciliation", "0",
    S_MAXABS(ref('EBIT_Analysis','ros_diff')), 0.05, "Calculation matches Porsche")
chk("C29", "Own EBIT change = Porsche bridge buckets (periods outside tolerance)", "Porsche reconciliation", "0",
    S_COUNT_GT(ref('EBIT_Analysis','pbdiff','F'), ref('EBIT_Analysis','pbtol','F')), 0, "Our bridge total agrees with Porsche")
chk("C30", "Capitalisation rate = published (max abs. diff, pp)", "Porsche reconciliation", "0",
    S_MAXABS(ref('RD_DA','crdiff')), 0.15, "Calculation matches Porsche")
chk("C31", "ASP = published ASP (max abs. diff, EUR k; published rounded to 1k)", "Porsche reconciliation", "0",
    S_MAXABS(ref('Commercial','aspd')), 0.5, "Porsche ASP definition reproduced")
chk("C32", "NCF margin = published (max abs. diff, pp)", "Porsche reconciliation", "0",
    S_MAXABS(ref('Cash_Crosscheck','ncfm_d')), 0.05, "Calculation matches Porsche")
chk("C33", "EBITDA margin = published (max abs. diff, pp)", "Porsche reconciliation", "0",
    S_MAXABS(ref('Cash_Crosscheck','ebm_d')), 0.05, "Calculation matches Porsche")
chk("C34", "Volume + ASP effect = revenue change", "Identity", "0", S_ABS(ref('Commercial','vchk','F')), 0.001, "Revenue bridge exact")
chk("C35", "Deliveries by model and by region = total deliveries", "Identity", "0",
    S_ABS(ref('Commercial','m_chk')) + "+" + S_ABS(ref('Commercial','r_chk')), 0, "Mix tables complete")
chk("C36", "NCF = operating + investing cash flow (AF-12 rounding)", "Identity", "0", S_ABS(ref('Cash_Crosscheck','ncfchk')), 0.5, "Cash data consistent")
chk("C37", "Guidance grid: 27 combinations and FY EBIT = H1 2026 EBIT + H2 EBIT", "Identity", "0",
    f"ABS(COUNTA({g}!A{GRID+1}:A{GRID_L})-27)+" + "+".join(f"ABS(N({g}!I{r})-N({A('h1ebit').replace('$C$', g + '!$C$')})-N({g}!K{r}))" for r in range(GRID + 1, GRID_L + 1)), 0.001, "Scenario arithmetic consistent")
chk("C38", "Labelling: every guidance row says 'not a forecast'", "Labelling", "0",
    f"ABS(COUNTA({g}!B{GRID+1}:B{GRID_L})-27)+ABS(COUNTA({g}!B{AG+1}:B{AG_L})-9)+" + S_NOT_EQ(f"{g}!B{GRID+1}:B{GRID_L}", SCEN_LABEL) + "+" + S_NOT_EQ(f"{g}!B{AG+1}:B{AG_L}", SCEN_LABEL), 0, "Scenarios cannot be mistaken for forecasts")
chk("C39", "Labelling: L3 section flagged as supplementary hybrid sensitivity", "Labelling", "0",
    f'IF(EBIT_Analysis!A{REG[("EBIT_Analysis","daclean")]-2}="{L3_HEADING}",0,1)', 0, "L3 cannot be mistaken for the headline KPI")
chk("C40", "SQL data-quality checks (Phase 3 + 4): FAIL count", "SQL", "0", f"SUM(Audit_Register!D{DQ_TOP}:D{DQ_BOT})", 0, "Upstream database validated")
chk("C41", "Data_Long: every (period_code, var_id) pair is unique (max count)", "Lookup integrity", "1", f"MAX(Data_Long!K{DL_FIRST}:K{DL_LAST})", 0, "Two-key SUMIFS can never add two rows together")
chk("C42", "Data_Panel round-trip: each Data_Long value read back from the panel grid (sum of abs. differences)", "Lookup integrity", "0",
    f"SUM(Data_Long!M{DL_FIRST}:M{DL_LAST})", 0.000001, "Every panel cell holds the value of its own period and variable")
# --- anchor values: the numbers DISPLAYED on KPI_Overview vs the SQL export (literal snapshot, blue)
_fl = {(r[0], r[1]): float(r[3]) for r in rd("fact_long.csv")[1:]}
_so = {r[0]: r for r in rd("sqlout.csv")[1:]}; _sh = rd("sqlout.csv")[0]
_sv = lambda p, col: float(_so[p][_sh.index(col)])
_kpi = WS["KPI_Overview"]
def _krow(label):
    for r in range(1, 60):
        if _kpi.cell(row=r, column=2).value == label: return r
    raise KeyError(label)
_pcol = {p: PC[k] for k, p in enumerate(P)}
ANCH = [
 ("A01", "2023-H1 Group revenue (EUR m)", "Group sales revenue", "2023-H1", _fl[("2023-H1", "G01")], 1, 0.000001, "core.fact_financial G01"),
 ("A02", "2023-H1 Group EBIT (EUR m)", "L0 Reported EBIT", "2023-H1", _sv("2023-H1", "ebit_l0_reported"), 1, 0.000001, "v_ebit_adjusted"),
 ("A03", "2025-H1 Group EBIT (EUR m)", "L0 Reported EBIT", "2025-H1", _sv("2025-H1", "ebit_l0_reported"), 1, 0.000001, "v_ebit_adjusted"),
 ("A04", "2026-H1 Group EBIT (EUR m)", "L0 Reported EBIT", "2026-H1", _sv("2026-H1", "ebit_l0_reported"), 1, 0.000001, "v_ebit_adjusted"),
 ("A05", "2026-H1 reported margin (%, SQL rounded 0.01)", "L0 Reported margin", "2026-H1", _sv("2026-H1", "margin_l0_reported_pct"), 100, 0.005, "v_ebit_adjusted"),
 ("A06", "2026-H1 L1 EBIT (EUR m)", "L1 EBIT excl. realignment & battery", "2026-H1", _sv("2026-H1", "ebit_l1_ex_realignment"), 1, 0.000001, "v_ebit_adjusted"),
 ("A07", "2026-H1 L1 margin (%, SQL rounded 0.01)", "L1 margin", "2026-H1", _sv("2026-H1", "margin_l1_pct"), 100, 0.005, "v_ebit_adjusted"),
 ("A08", "2026-H1 ASP (EUR k, SQL rounded 0.01)", "ASP (automotive revenue per vehicle sold)", "2026-H1", _sv("2026-H1", "asp_calc_eur_k"), 1, 0.005, "v_asp_volume"),
]
for k, (aid, text, lab, p, val, mult, tol, view) in enumerate(ANCH):
    cell = f"KPI_Overview!{_pcol[p]}{_krow(lab)}"
    chk(f"C{43+k}", f"Anchor {aid}: {text} displayed on KPI_Overview = SQL {view} ({val:,.6f})", "Anchor (displayed vs SQL)",
        repr(val), f"{cell}*{mult}", tol, f"Displayed value {cell} equals the SQL export")
for i, (cid, text, cat_, expected, actual, tol, proves) in enumerate(CH):
    rr = 7 + i
    ws.cell(row=rr, column=1, value=cid).font = F_N
    ws.cell(row=rr, column=2, value=text).font = F_N
    ws.cell(row=rr, column=3, value=cat_).font = F_N
    ws.cell(row=rr, column=4, value="=" + expected).font = F_N
    ws.cell(row=rr, column=5, value="=" + actual).font = F_N
    ws.cell(row=rr, column=6, value=f"=E{rr}-D{rr}").font = F_N
    c = ws.cell(row=rr, column=7, value=tol); c.font = F_IN
    ws.cell(row=rr, column=8, value=f'=IF(ABS(F{rr})<=G{rr},"PASS","FAIL")').font = F_B
    ws.cell(row=rr, column=9, value=proves).font = F_SUB
    for col in "DEFG": ws[f"{col}{rr}"].number_format = '#,##0.0000'
    for col in "ABCDEFGHI": ws[f"{col}{rr}"].border = BD
CH_L = 7 + len(CH) - 1

STATUS_F = (f'IF(COUNTIF(Audit_Checks!$H$7:$H${CH_L},"FAIL")=0,"ALL "&COUNTA(Audit_Checks!$H$7:$H${CH_L})&" CHECKS PASS",'
            f'COUNTIF(Audit_Checks!$H$7:$H${CH_L},"FAIL")&" OF "&COUNTA(Audit_Checks!$H$7:$H${CH_L})&" CHECK(S) FAIL - see Audit_Checks")')
ws["C4"] = "=" + STATUS_F
ws["C4"].font = Font(name=AR, size=11, bold=True, color="006100"); ws.merge_cells("C4:F4")
ws.conditional_formatting.add(f"H7:H{CH_L}", FormulaRule(formula=[f'H7="FAIL"'], fill=PatternFill("solid", start_color="FFC7CE")))
ws.conditional_formatting.add(f"H7:H{CH_L}", FormulaRule(formula=[f'H7="PASS"'], fill=PatternFill("solid", start_color="C6EFCE")))
ws.freeze_panes = "A7"

# ================================================================ README
ws = WS["README"]
title(ws, "Porsche Company Lens - 2026 Margin Recovery: business-analysis workbook",
      "Phase 5 (Excel). Built from the validated PostgreSQL database porsche_lens (Phases 3-4). Prepared 2026-09-28.",
      {"A": 26, "B": 110})
lines = [
 ("Question", "Porsche's reported margin rose from 5.5% to 7.8% in H1 2026. After removing the exceptional items Porsche itself discloses, "
              "and isolating non-cash effects from R&D capitalisation and depreciation, is underlying profitability actually improving, "
              "and what does FY2026 guidance imply for H2 2026?"),
 ("Purpose", "A transparent, formula-driven model that a business team can audit: every number traces back to a Porsche document and page."),
 ("Model status", "=MODEL_STATUS_PLACEHOLDER"),
 ("", ""),
 ("SHEET MAP", "What each sheet answers"),
 ("KPI_Overview", "What happened? Headline KPIs by half-year, H1 2025 vs H1 2026, formula-driven readings, two charts."),
 ("EBIT_Analysis", "Is underlying profitability improving? Reported (L0) vs adjusted (L1, L2) EBIT and margin; own YoY bridge in EUR m and pp; reconciliation to Porsche's bridge. L3 in a grey supplementary box."),
 ("RD_DA", "Are non-cash effects flattering EBIT? Total R&D costs, capitalisation rate, P&L charge, counterfactual; D&A cleaned of impairments."),
 ("Commercial", "Is value compensating for volume? ASP (Porsche definition), vehicle sales, volume vs ASP effect, model and regional mix."),
 ("Cash_Crosscheck", "Does cash confirm the profit picture? EBIT/EBITDA vs net cash flow, published-margin reconciliation, simple cash proxy."),
 ("Guidance_Scenarios", "What must H2 2026 look like to meet guidance? Scenario selector + all 27 combinations + automotive view. Scenarios, NOT forecasts."),
 ("Audit_Checks", f"Can the numbers be trusted? {len(CH)} formula checks: imports, lookup integrity (unique keys + panel round-trip), Excel vs SQL reconciliation, identities, Porsche reconciliation, labelling, and 8 anchor values displayed on KPI_Overview vs the SQL export."),
 ("Audit_Register", "Where does it come from? Source register, assumptions A-01..A-09, known data gaps, audit findings, SQL view catalogue, SQL check results."),
 ("Data_Panel", "Half-year panel (formula pivot of Data_Long) with a FACT/DERIVED label grid; derived cells shaded."),
 ("Data_Long / Data_Mix", "Controlled imports from SQL (core.fact_financial, deliveries). The only place where Porsche values are typed in."),
 ("SQL_Outputs", "Mart view results, used ONLY by Audit_Checks to prove that Excel and SQL agree."),
 ("", ""),
 ("CONVENTIONS", ""),
 ("Colours", "Blue = imported or hard-coded input. Black = formula. Green = link to another sheet. Yellow fill = headline KPI. Grey fill = supplementary sensitivity. Orange band = scenario area."),
 ("Units", "Money in EUR m. Margins stored as fractions and shown as %. 'pp' = percentage points. ASP in EUR k. Volumes in units."),
 ("Signs", "Exceptional items and tariffs are positive numbers = expense; adding them back raises EBIT. Bridge effects: + = raises EBIT."),
 ("Periods", "Half-years 2023-H1 .. 2026-H1. Porsche publishes H1 and FY; H2 = FY - H1 (DERIVED in SQL). Year-on-year always compares the same half of the prior year."),
 ("FACT / DERIVED", "FACT = published by Porsche. DERIVED = calculated from published values (H2, identities, FY2025 one-off total). See Data_Panel label grid and Data_Long source_ref."),
 ("", ""),
 ("PROFIT LEVELS", ""),
 ("L0 (headline)", "Reported group EBIT."),
 ("L1 (headline)", "L0 + strategic realignment & battery items as disclosed by Porsche (net of related provision releases)."),
 ("L2 (headline)", "L1 + US import tariffs."),
 ("L3 (supplementary)", "Supplementary HYBRID SENSITIVITY only: L1 + clean automotive D&A - automotive capitalised development costs. Not underlying group profitability; never used as headline KPI."),
 ("Residual", "'Residual / other EBIT drivers' in the bridge is NOT operating performance: it mixes price, mix, volume, total R&D cost changes, material/supplier costs, SG&A, other operating result, FX, consolidation and one-off rounding."),
 ("", ""),
 ("KEY ASSUMPTIONS", "Full register: Audit_Register section 2."),
 ("A-01 / A-02", "No realignment items disclosed for 2023-2024 and no US tariffs before April 2025 -> treated as 0."),
 ("A-03 / A-04", "Undisclosed impairments (H2 2023, H2 2024) -> 0; H1 2026 impairments (EUR 61m) assumed inside D&A (alternative shown)."),
 ("A-05..A-08", "Guidance: automotive revenue share, extraordinary-expense basis, H2 tariffs = H1, range ends combined independently."),
 ("A-09", "Derived H2 values inherit the EUR 0.1bn rounding of one-off disclosures (rounding band shown)."),
 ("", ""),
 ("LINEAGE", "Porsche documents (S001-S027) -> PCL_02_Manual_Inputs.xlsx + fact sheet S025 -> PostgreSQL core.fact_financial -> mart views -> Data_Long (this workbook) -> analysis sheets. "
             "Refresh: re-export fact_long / mix / SQL outputs and replace the blue import rows; Audit_Checks must show ALL PASS."),
 ("Not in scope", "No forecasts, no internal Porsche data, no estimates for missing values. Interpretation and recommendations follow in Phases 6-9."),
]
for i, (a, b) in enumerate(lines):
    rr = 4 + i
    ca = ws.cell(row=rr, column=1, value=a); ca.font = F_B if a.isupper() or a in ("Question", "Purpose", "Model status") else F_N
    cb = ws.cell(row=rr, column=2, value=b); cb.font = F_B if a.isupper() else (F_LINK if str(b).startswith("=") else F_N); cb.alignment = WR
    if a.isupper() and a:
        for col in (1, 2): ws.cell(row=rr, column=col).fill = FILL_SEC
    if a in ("L3 (supplementary)",):
        for col in (1, 2): ws.cell(row=rr, column=col).fill = FILL_SUPP
    if a == "Guidance_Scenarios":
        for col in (1, 2): ws.cell(row=rr, column=col).fill = FILL_SCEN

# tab colours
for n, colr in (("README", "1F3864"), ("KPI_Overview", "FFC000"), ("EBIT_Analysis", "4472C4"), ("RD_DA", "4472C4"), ("Commercial", "4472C4"),
                ("Cash_Crosscheck", "4472C4"), ("Guidance_Scenarios", "ED7D31"), ("Audit_Checks", "70AD47"), ("Audit_Register", "70AD47"),
                ("Data_Panel", "A5A5A5"), ("Data_Long", "A5A5A5"), ("Data_Mix", "A5A5A5"), ("SQL_Outputs", "A5A5A5")):
    WS[n].sheet_properties.tabColor = colr
for n in ORDER:
    WS[n].sheet_view.zoomScale = 90
for _ws in wb:
    for _row in _ws.iter_rows():
        for _c in _row:
            if isinstance(_c.value, str) and "MODEL_STATUS_PLACEHOLDER" in _c.value:
                _c.value = _c.value.replace("MODEL_STATUS_PLACEHOLDER", STATUS_F)
from openpyxl.worksheet.properties import PageSetupProperties
for _ws in wb:
    _ws.page_setup.orientation = "landscape"; _ws.page_setup.paperSize = _ws.PAPERSIZE_A4
    _ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    _ws.page_setup.fitToWidth = 1; _ws.page_setup.fitToHeight = 0
    _ws.print_options.gridLines = False
wb.calculation.fullCalcOnLoad = True
import json as _json
_json.dump({"REG": {f"{k[0]}|{k[1]}": v for k, v in REG.items()}, "PROW": PROW, "SQLROW": SQLROW,
            "recon": [(r[0], r[2], r[3], r[4]) for r in recon], "P": P, "PC": PC},
           open(os.path.splitext(OUT)[0] + "_map.json", "w"), indent=1)
wb.save(OUT)
print("saved", OUT, "checks:", len(CH))
